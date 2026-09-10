from __future__ import annotations

from openpyxl import load_workbook

from src.export.consolidation import TOP_COLUMNS, balanced_pending, priority, retrieval_class, write_workbook


def test_priority_is_transparent_and_never_changes_review_status() -> None:
    row = {
        "Candidate ID": "C-1",
        "Retrieval Tier": "PRIMARY",
        "Retrieval Signal": "Strong lexical signal 'x'",
        "Stance Hint": "POSSIBLE_ABUSE",
        "Review Status": "PENDING",
        "Source Item ID": "ITEM-1",
    }
    klass = retrieval_class(row, "")
    score, basis = priority(row, klass, True)

    assert klass == "STRONG_LEXICAL"
    assert 0 <= score <= 100
    assert "strong_lexical" in basis
    assert row["Review Status"] == "PENDING"


def test_balanced_pending_does_not_drop_or_duplicate_candidates() -> None:
    rows = [
        {"Candidate ID": "A", "Source URL": "source-1", "Review Priority Score": "99"},
        {"Candidate ID": "B", "Source URL": "source-1", "Review Priority Score": "98"},
        {"Candidate ID": "C", "Source URL": "source-2", "Review Priority Score": "90"},
        {"Candidate ID": "D", "Source URL": "source-3", "Review Priority Score": "80"},
    ]
    selected = balanced_pending(rows, 4)

    assert {row["Candidate ID"] for row in selected} == {"A", "B", "C", "D"}
    assert [int(row["Balanced Review Rank"]) for row in selected] == [1, 2, 3, 4]


def test_top_workbook_is_concise_and_source_link_is_clickable(tmp_path) -> None:
    path = tmp_path / "top.xlsx"
    row = {column: "" for column in TOP_COLUMNS}
    row.update({
        "Balanced Review Rank": "1",
        "Candidate ID": "C-1",
        "Example": "Exact real comment",
        "Source Link": "https://research-fixture.invalid/discussion",
        "Source Item ID": "ITEM-1",
        "Review Status": "PENDING",
        "Review Priority Score": "90",
    })
    write_workbook(path, {"GEN-1": [row]}, top=True)

    workbook = load_workbook(path)
    sheet = workbook["GEN-1"]
    headers = [cell.value for cell in sheet[1]]
    link_cell = sheet.cell(row=2, column=headers.index("Source Link") + 1)
    assert headers == TOP_COLUMNS
    assert link_cell.hyperlink.target == row["Source Link"]
    workbook.close()

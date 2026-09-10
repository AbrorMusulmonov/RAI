from copy import deepcopy
import json
import pytest
from openpyxl import Workbook,load_workbook
from src.audit.final50 import AUDIT, ROOT, CATEGORIES, HUMAN, queues, all_raw, evidence_index
from src.export.final_human_review import COLUMNS,STATUSES,excel_text,write_main,ordered,assert_no_unmerged_edits

def test_existing_human_rows_and_healthy_category_unchanged():
    baseline=json.loads((AUDIT/'baseline.json').read_text(encoding='utf-8'))
    current=queues()
    for c,rows in baseline['queues'].items():
        byid={r['Candidate ID']:r for r in current[c]}
        for r in rows:
            if r['Review Status'] in HUMAN:
                assert r['Candidate ID'] in byid
                assert all(byid[r['Candidate ID']].get(k)==v for k,v in r.items())
    # Only audited PENDING noise may leave a former completed category.
    archived={r['Candidate ID'] for r in __import__('src.processing.category',fromlist=['csv_rows']).csv_rows(ROOT/'data/candidates/GEN-4.final50_quality_excluded.csv')}
    assert {r['Candidate ID'] for r in baseline['queues']['GEN-4']} <= {r['Candidate ID'] for r in current['GEN-4']}|archived
    assert json.loads((ROOT/'data/reviewed/final_ids.json').read_text(encoding='utf-8'))==baseline['final_ids']

def test_final_workbook_exact_canonical_rows_and_controls():
    q=queues();w=load_workbook(ROOT/'data/exports/FINAL_HUMAN_REVIEW.xlsx')
    assert set(w.sheetnames)==set(CATEGORIES)|{'DASHBOARD','INSTRUCTIONS'}
    for c,rows in q.items():
        ws=w[c];assert [x.value for x in ws[1]]==COLUMNS
        actual={r[0]:dict(zip(COLUMNS,r)) for r in ws.iter_rows(min_row=2,values_only=True) if r[0]}
        assert set(actual)=={r['Candidate ID'] for r in rows}
        for r in rows:
            a=actual[r['Candidate ID']]
            assert a['Example']==r['Example']
            assert a['Review Status']==r['Review Status']
            assert (a['Reviewer Notes'] or '')==(r.get('Reviewer Notes') or '')
            assert a['Source Link']==(r.get('Content URL') or r['Source URL'])
        assert not ws.merged_cells
        assert ws.freeze_panes
        assert ws.data_validations.dataValidation[0].formula1=='"'+','.join(STATUSES)+'"'
        for row in ws.iter_rows(min_row=2):
            assert row[3].hyperlink.target==row[3].value
            assert row[1].data_type=='s'
    assert w['DASHBOARD']['E2'].value.startswith('=COUNTIF(')
    assert w['DASHBOARD']['M2'].value=='=MAX(0,50-E2)'
    w.close()

def test_new_rows_pending_real_and_no_invented_permalink():
    b=json.loads((AUDIT/'baseline.json').read_text(encoding='utf-8'));q=queues();raw,_=all_raw();index=evidence_index(raw)
    for c,rows in q.items():
        old={r['Candidate ID'] for r in b['queues'][c]}
        for r in rows:
            ev=index[(r['Source URL'],r['Source Item ID'],r['Example'])]
            assert ev
            assert r.get('Human-readable Provenance') or r.get('Provenance')
            if r['Candidate ID'] not in old:assert r['Review Status']=='PENDING'
            if r.get('Content URL'):assert any(e.get('direct_permalink_available') and e.get('content_url')==r['Content URL'] for e in ev)

def test_public_text_cannot_become_excel_formula():
    w=Workbook();cell=w.active['A1'];excel_text(cell,'=1+1');assert cell.data_type=='s';assert cell.value=='=1+1'

def test_export_refuses_unmerged_excel_decisions_or_notes(tmp_path):
    q=queues();c=next(c for c in CATEGORIES if q[c]);r=deepcopy(q[c][0]);r['Review Status']='PENDING'
    w=Workbook();ws=w.active;ws.title=c;ws.append(COLUMNS)
    ws.append([r.get(k,'') if k!='Reviewer Notes' else 'human note not yet merged' for k in COLUMNS]);p=tmp_path/'review.xlsx';w.save(p)
    with pytest.raises(RuntimeError,match='Refusing overwrite'):assert_no_unmerged_edits(p,q)

def test_aziza_export_only_existing_accepts():
    q=queues();w=load_workbook(ROOT/'data/exports/aziza_accepted.xlsx')
    for c,rows in q.items():
        accepted=[r['Example'] for r in rows if r['Review Status']=='ACCEPT']
        assert [r[1] for r in w[c].iter_rows(min_row=2,values_only=True)]==accepted
    w.close()


def test_dashboard_formulas_reference_exact_category_ranges():
    q=queues();w=load_workbook(ROOT/'data/exports/FINAL_HUMAN_REVIEW.xlsx')
    for i,c in enumerate(CATEGORIES,2):
        end=max(2,len(q[c])+1);ref=f"'{c}'!";d=w['DASHBOARD']
        assert d.cell(i,2).value==f'=COUNTA({ref}$A$2:$A${end})'
        for col,status in enumerate(('ACCEPT','REJECT','UNSURE','MOVE_TO_OTHER_CATEGORY','PENDING'),5):
            assert d.cell(i,col).value==f'=COUNTIF({ref}$M$2:$M${end},"{status}")'
        assert d.cell(i,12).value==f'=E{i}+I{i}*K{i}'
        assert d.cell(i,13).value==f'=MAX(0,50-E{i})'
    assert w.calculation.calcMode=='auto'
    assert w.calculation.fullCalcOnLoad
    w.close()


def test_raw_prefixes_exclusions_and_initial_recovery_preserved():
    import hashlib
    from src.processing.category import csv_rows
    from src.audit.final50_recovery import PICKS,REMOVE
    b=json.loads((AUDIT/'baseline.json').read_text(encoding='utf-8'));q=queues()
    for rel,meta in b['files'].items():
        if rel.startswith('data/raw/'):
            with (ROOT/rel).open('rb') as f:assert hashlib.sha256(f.read(meta['size'])).hexdigest()==meta['sha256']
    s=json.loads((AUDIT/'screen.json').read_text(encoding='utf-8'))
    for c in CATEGORIES:
        active={r['Candidate ID'] for r in q[c]}
        archived={r['Candidate ID'] for r in csv_rows(ROOT/f'data/candidates/{c}.final50_quality_excluded.csv')}
        assert not active&archived
        assert set(REMOVE.get(c,{}))<=archived
        for indices in PICKS.get(c,{}).values():
            assert all(s['categories'][c][i]['candidate_id'] in active|archived for i in indices)


def test_new_ids_and_occurrences_are_deterministic():
    from src.utils.io import read_jsonl
    from src.processing.category import stable_candidate_id,lexical
    from src.processing.gen1 import stable_candidate_id as gen1_id
    from src.utils.text import normalized_for_comparison
    b=json.loads((AUDIT/'baseline.json').read_text(encoding='utf-8'))
    for c,rows in queues().items():
        old={r['Candidate ID'] for r in b['queues'][c]}
        occurrence={(r.get('candidate_id'),r.get('raw_record_id')) for r in read_jsonl(ROOT/f'data/candidates/{c}.occurrences.jsonl')}
        for r in rows:
            if r['Candidate ID'] in old:continue
            expected=gen1_id(normalized_for_comparison(r['Example'])) if c=='GEN-1' else stable_candidate_id(c,lexical(r['Example']))
            assert r['Candidate ID']==expected
            assert (expected,r['Final50 Raw ID']) in occurrence
            assert r['Review Status']=='PENDING'


def test_active_derived_views_follow_canonical_review_state():
    from src.processing.category import csv_rows
    for c,rows in queues().items():
        for suffix,tier in [('.csv','PRIMARY'),('.secondary_review.csv','SECONDARY')]:
            actual=csv_rows(ROOT/f'data/candidates/{c}{suffix}')
            expected=[r for r in rows if r.get('Retrieval Tier')==tier]
            keys=('Candidate ID','Example','Review Status','Reviewer Notes')
            assert {tuple(r.get(k,'') for k in keys) for r in actual}=={tuple(r.get(k,'') for k in keys) for r in expected}


def test_quality_companion_notebook_executes_read_only():
    notebook=json.loads((ROOT/'notebooks/final50_quality_audit.ipynb').read_text(encoding='utf-8'))
    state={}
    for cell in notebook['cells']:
        if cell['cell_type']=='code':exec(compile(''.join(cell['source']),'final50_quality_audit.ipynb','exec'),state)

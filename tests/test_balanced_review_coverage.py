"""Coverage is independent of ACCEPT; no private examples are embedded in tests."""
from copy import deepcopy
import pytest
from openpyxl import load_workbook
from src.audit.final50 import CATEGORIES
from src.export.final_human_review import write_main, COLUMNS, DASH_COLUMNS


@pytest.fixture
def review_inputs():
    """Reserved-domain fixtures, not collected research examples or decisions."""
    current={c:[] for c in CATEGORIES}
    for i,status in enumerate(('PENDING','ACCEPT','REJECT','UNSURE','MOVE_TO_OTHER_CATEGORY')):
        current['GEN-1'].append({
            'Candidate ID':f'TEST-ONLY-{i}',
            'Example':f'[TEST ONLY: literal cell {i}]',
            'Source URL':'https://example.invalid/test-only',
            'Source Item ID':f'TEST-ITEM-{i}',
            'Retrieval Tier':'PRIMARY' if i%2==0 else 'SECONDARY',
            'Review Status':status,
            'Reviewer Notes':f'TEST ONLY: preserved note {i}',
        })
    matrix={c:{'Distinct Sources':int(bool(rows)),
               'Estimated Review Precision':0.5,
               'Readiness Status':'TEST ONLY'} for c,rows in current.items()}
    counts={c:len(rows) for c,rows in current.items()}
    # Audited coverage need not equal the full queue or any decision count.
    counts['GEN-1']=3
    return current,matrix,counts


def test_audited_candidate_coverage_is_not_accept_count(tmp_path,review_inputs):
    q,matrix,counts=review_inputs
    path=tmp_path/'review.xlsx'
    original=deepcopy(q)
    write_main(q,matrix,path,verified_active=counts)
    assert q==original
    w=load_workbook(path)
    d=w['DASHBOARD']
    assert d['M1'].value=='Human ACCEPT Gap to 50'
    assert d['K1'].value=='Prior Audit Precision (not revalidated)'
    assert d['L1'].value=='Illustrative Yield (prior precision)'
    assert d['P1'].value=='Gap to 50'
    assert d['Q1'].value=='Target Status'
    for i,c in enumerate(CATEGORIES,2):
        assert d.cell(i,15).value==counts[c]
        assert d.cell(i,16).value==f'=MAX(0,50-O{i})'
        assert d.cell(i,17).value==f'=IF(O{i}<50,"INCOMPLETE",IF(O{i}<60,"MINIMUM TARGET MET",IF(O{i}<=70,"TARGET COMPLETE","ABOVE TARGET")))'
        assert [x.value for x in w[c][1]]==COLUMNS
        assert d.cell(i,13).value==f'=MAX(0,50-E{i})'
        actual={r[0]:dict(zip(COLUMNS,r)) for r in w[c].iter_rows(min_row=2,values_only=True) if r[0]}
        for row in q[c]:
            exported=actual[row['Candidate ID']]
            assert exported['Example']==row['Example']
            assert exported['Source Item ID']==row['Source Item ID']
            assert exported['Source Link']==row['Source URL']
            assert exported['Review Status']==row['Review Status']
            assert exported['Reviewer Notes']==row['Reviewer Notes']
    assert d.auto_filter.ref=='A1:Q17'
    w.close()


def test_coverage_cannot_exceed_real_queue_or_use_invalid_counts(tmp_path,review_inputs):
    q,matrix,counts=review_inputs
    for invalid in (-1,True,1.5,len(q['GEN-1'])+1):
        bad={**counts,'GEN-1':invalid}
        with pytest.raises(ValueError,match='Invalid audited active coverage'):
            write_main(q,matrix,tmp_path/'bad.xlsx',verified_active=bad)
    assert not (tmp_path/'bad.xlsx').exists()


def test_default_dashboard_remains_backward_compatible(tmp_path,review_inputs):
    q,matrix,_=review_inputs
    path=tmp_path/'legacy.xlsx'
    write_main(q,matrix,path)
    w=load_workbook(path)
    d=w['DASHBOARD']
    assert [cell.value for cell in d[1]]==DASH_COLUMNS
    assert d.auto_filter.ref=='A1:N17'
    assert d['M2'].value=='=MAX(0,50-E2)'
    assert d['L2'].value=='=E2+I2*K2'
    w.close()


def test_coverage_export_refuses_unmerged_human_edits(tmp_path,review_inputs):
    q,matrix,counts=review_inputs
    path=tmp_path/'review.xlsx'
    write_main(q,matrix,path,verified_active=counts)
    w=load_workbook(path)
    ws=w['GEN-1']
    pending_row=next(i for i in range(2,ws.max_row+1) if ws.cell(i,1).value=='TEST-ONLY-0')
    ws.cell(pending_row,13,'ACCEPT')
    w.save(path)
    w.close()
    previous=path.read_bytes()
    with pytest.raises(RuntimeError,match='Refusing overwrite'):
        write_main(q,matrix,path,verified_active=counts)
    assert path.read_bytes()==previous
    assert q['GEN-1'][0]['Review Status']=='PENDING'


def test_instruction_freeze_has_no_duplicate_or_stale_panes(tmp_path,review_inputs):
    q,matrix,counts=review_inputs
    path=tmp_path/'review.xlsx'
    write_main(q,matrix,path,verified_active=counts)
    w=load_workbook(path)
    instructions=w['INSTRUCTIONS']
    assert instructions.freeze_panes=='A2'
    assert instructions.sheet_view.pane.activePane=='bottomLeft'
    assert [s.pane for s in instructions.sheet_view.selection]==['bottomLeft']
    for ws in w:
        panes=[s.pane for s in ws.sheet_view.selection]
        assert len(panes)==len(set(panes))
    w.close()


def test_table_filters_do_not_overlap_sheet_filters(tmp_path,review_inputs):
    q,matrix,_=review_inputs
    path=tmp_path/'review.xlsx'
    write_main(q,matrix,path)
    w=load_workbook(path)
    for category,rows in q.items():
        ws=w[category]
        if rows:
            assert ws.auto_filter.ref is None
            assert len(ws.tables)==1
            table=next(iter(ws.tables.values()))
            assert table.ref==f'A1:N{len(rows)+1}'
            assert table.autoFilter.ref==table.ref
        else:
            assert not ws.tables
            assert ws.auto_filter.ref=='A1:N1'
    w.close()

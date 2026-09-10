from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from src.utils.io import PROJECT_ROOT


CATEGORIES = ("GEN-1", "REG-1", "REG-3", "REG-4")
HUMAN = {"ACCEPT", "REJECT", "UNSURE", "MOVE_TO_OTHER_CATEGORY"}


def rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def test_gap_fill_preserves_human_decision_semantics() -> None:
    baseline = json.loads((PROJECT_ROOT / "data/audits/gap_fill_a_baseline.json").read_text(encoding="utf-8"))
    current: dict[str, dict[str, str]] = {}
    for category in CATEGORIES:
        for row in rows(PROJECT_ROOT / f"data/reviewed/{category}.review_queue.csv"):
            if row.get("Review Status") in HUMAN:
                current[f"{category}:{row['Candidate ID']}"] = row
    assert set(current) == set(baseline["human_decisions"])
    for key, expected in baseline["human_decisions"].items():
        assert {field: current[key].get(field, "") for field in expected} == expected


def test_gap_fill_new_candidates_are_pending_and_exclusions_are_auditable() -> None:
    baseline = json.loads((PROJECT_ROOT / "data/audits/gap_fill_a_baseline.json").read_text(encoding="utf-8"))
    cleanup = json.loads((PROJECT_ROOT / "config/gap_fill_a_cleanup.json").read_text(encoding="utf-8"))
    for category in CATEGORIES:
        queue = rows(PROJECT_ROOT / f"data/reviewed/{category}.review_queue.csv")
        old_ids = set(baseline["queues"][category]["candidate_ids"])
        assert all(row["Review Status"] == "PENDING" for row in queue if row["Candidate ID"] not in old_ids)
        excluded = rows(PROJECT_ROOT / f"data/candidates/{category}.gap_fill_quality_excluded.csv")
        assert {row["Candidate ID"] for row in excluded} == set(cleanup["categories"][category])
        assert all(row["Review Status"] == "PENDING" for row in excluded)


def test_gap_fill_workbook_is_pending_only_with_real_links() -> None:
    workbook = load_workbook(PROJECT_ROOT / "data/exports/gap_fill_review_batch_A.xlsx")
    assert workbook.sheetnames == list(CATEGORIES)
    for category in CATEGORIES:
        sheet = workbook[category]
        headers = [cell.value for cell in sheet[1]]
        status_column = headers.index("Review Status") + 1
        link_column = headers.index("Source Link") + 1
        for row_number in range(2, sheet.max_row + 1):
            assert sheet.cell(row_number, status_column).value == "PENDING"
            link = sheet.cell(row_number, link_column)
            assert str(link.value).startswith("https://")
            assert link.hyperlink is not None


def test_gap_fill_final_integrity_audit_passes() -> None:
    audit = json.loads((PROJECT_ROOT / "data/audits/gap_fill_a_final.json").read_text(encoding="utf-8"))
    assert audit["passed"] is True
    assert audit["errors"] == []
    assert audit["totals"]["synthetic_examples_created"] == 0
    assert audit["totals"]["fabricated_urls_created"] == 0
    assert audit["totals"]["provenance_failures"] == 0
    assert audit["integrity"]["human_decisions_overwritten"] == 0

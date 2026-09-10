from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.export.__main__ import source_is_verified
from src.processing.category import lexical
from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


REQUIRED_NEW_QUEUE_FIELDS = {
    "Candidate ID", "Example", "Context", "Source Platform", "Source Name",
    "Source URL", "Content URL", "Source Item ID", "Retrieval Tier",
    "Retrieval Signal", "Stance Hint", "Stance Evidence", "Suggested Category",
    "Possible Compound Category", "Review Status", "Reviewer Notes",
}
CANONICAL_TERM_KEYS = {
    "strong_terms", "contextual_terms", "discovery_queries", "weak_terms",
    "spelling_variants", "target_terms", "harm_terms", "standalone_strong_terms",
    "compound_patterns",
}


def csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.exists():
        return [], []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def audit() -> dict[str, Any]:
    taxonomy = read_json(PROJECT_ROOT / "config" / "taxonomy.json")
    codes = [item["code"] for item in taxonomy["subcategories"]]
    new_codes = [code for code in codes if code != "GEN-1"]
    registry = read_json(PROJECT_ROOT / "config" / "sources.json")
    sources_by_url = {str(source.get("url") or ""): source for source in registry["sources"]}
    terms = read_json(PROJECT_ROOT / "config" / "search_terms.json")["categories"]

    critical: list[str] = []
    warnings: list[str] = []
    raw_by_id: dict[str, dict[str, Any]] = {}
    raw_counts: Counter[str] = Counter()
    raw_unique_by_category: dict[str, set[str]] = defaultdict(set)
    for path in (PROJECT_ROOT / "data" / "raw").rglob("*.jsonl"):
        if path.name.endswith(".provenance.jsonl"):
            continue
        category = path.stem
        for record in read_jsonl(path):
            raw_counts[category] += 1
            raw_id = str(record.get("raw_record_id") or "")
            if raw_id:
                raw_unique_by_category[category].add(raw_id)
                raw_by_id.setdefault(raw_id, record)

    occurrence_map: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for path in (PROJECT_ROOT / "data" / "candidates").glob("*.occurrences.jsonl"):
        for occurrence in read_jsonl(path):
            candidate_id = str(occurrence.get("candidate_id") or "")
            if candidate_id:
                occurrence_map[candidate_id].append(occurrence)

    category_rows: list[dict[str, Any]] = []
    all_review_sources: set[str] = set()
    all_candidate_ids: set[str] = set()
    normalized_locations: dict[str, list[str]] = defaultdict(list)
    total_verified = 0
    total_failures = 0
    for code in codes:
        metrics_path = PROJECT_ROOT / "data" / "candidates" / f"{code}.metrics.json"
        metrics = read_json(metrics_path) if metrics_path.exists() else {}
        fields, queue = csv_rows(PROJECT_ROOT / "data" / "reviewed" / f"{code}.review_queue.csv")
        if code in new_codes:
            missing_fields = REQUIRED_NEW_QUEUE_FIELDS - set(fields)
            if missing_fields:
                critical.append(f"{code}: review queue is missing columns {sorted(missing_fields)}")
            missing_terms = CANONICAL_TERM_KEYS - set(terms.get(code, {}))
            if missing_terms:
                critical.append(f"{code}: search bank is missing keys {sorted(missing_terms)}")

        statuses = Counter(row.get("Review Status") or "" for row in queue)
        if code in new_codes and set(statuses) - {"PENDING"}:
            critical.append(f"{code}: new-category queue contains non-PENDING decisions {dict(statuses)}")
        source_counts = Counter(row.get("Source URL") or "" for row in queue)
        all_review_sources.update(url for url in source_counts if url)
        verified = 0
        provenance_failures: list[str] = []
        direct = 0
        source_item_only = 0
        for row in queue:
            candidate_id = row.get("Candidate ID") or ""
            if not candidate_id or candidate_id in all_candidate_ids:
                critical.append(f"{code}: missing or cross-queue duplicate Candidate ID {candidate_id!r}")
            all_candidate_ids.add(candidate_id)
            normalized_locations[lexical(row.get("Example") or "")].append(f"{code}:{candidate_id}")
            source_url = row.get("Source URL") or ""
            source = sources_by_url.get(source_url)
            if not source:
                provenance_failures.append(f"{candidate_id}: source URL absent from registry")
            elif not source_is_verified(source):
                provenance_failures.append(f"{candidate_id}: source identity/reachability not verified")
            if not row.get("Source Item ID"):
                provenance_failures.append(f"{candidate_id}: source item ID missing")
            occurrences = occurrence_map.get(candidate_id, [])
            evidence = [raw_by_id.get(str(item.get("raw_record_id") or "")) for item in occurrences]
            evidence = [record for record in evidence if record]
            exact = [record for record in evidence if record.get("original_text") == row.get("Example")]
            if not exact:
                provenance_failures.append(f"{candidate_id}: no exact raw-text occurrence")
            elif not any(str(record.get("source_item_id") or "") == row.get("Source Item ID") for record in exact):
                provenance_failures.append(f"{candidate_id}: queue/raw item ID mismatch")
            content_url = row.get("Content URL") or ""
            if content_url:
                direct += 1
                if not any(str(record.get("content_url") or "") == content_url for record in exact):
                    provenance_failures.append(f"{candidate_id}: direct URL not returned in matching raw metadata")
            elif source_url and row.get("Source Item ID"):
                source_item_only += 1
            if not provenance_failures or not any(item.startswith(candidate_id + ":") for item in provenance_failures):
                verified += 1
        if provenance_failures:
            critical.extend(f"{code}: {failure}" for failure in provenance_failures)

        queue_tiers = Counter(row.get("Retrieval Tier") or "" for row in queue)
        primary = queue_tiers["PRIMARY"]
        secondary = queue_tiers["SECONDARY"]
        raw = raw_counts[code]
        unique_raw = len(raw_unique_by_category[code])
        metric_failures = int(metrics.get("unverified_provenance") or 0)
        total_failures += metric_failures + len(provenance_failures)
        total_verified += verified
        category_rows.append({
            "category": code,
            "raw": raw,
            "unique_raw": unique_raw,
            "primary": primary,
            "secondary": secondary,
            "total_review": len(queue),
            "distinct_sources": len([url for url in source_counts if url]),
            "candidates_per_source": dict(sorted(source_counts.items())),
            "verified_provenance": verified,
            "unverified": metric_failures + len(provenance_failures),
            "direct_permalink": direct,
            "source_url_plus_item_id": source_item_only,
            "human_accept": statuses.get("ACCEPT", 0),
            "human_reject": statuses.get("REJECT", 0),
            "human_unsure": statuses.get("UNSURE", 0),
            "human_move": statuses.get("MOVE_TO_OTHER_CATEGORY", 0),
            "pending": statuses.get("PENDING", 0),
            "stance_counts": dict(metrics.get("stance_counts") or {}),
            "near_duplicate_candidates_collapsed": int(metrics.get("near_duplicate_candidates_collapsed") or 0),
        })

    baseline_path = PROJECT_ROOT / "data" / "audits" / "GEN-1.protection_baseline.json"
    baseline = read_json(baseline_path)
    consolidation_final_path = PROJECT_ROOT / "data" / "audits" / "consolidation_final.json"
    consolidation_integrity = (
        read_json(consolidation_final_path).get("integrity_ok") is True
        if consolidation_final_path.exists() else False
    )
    gen1_protection: dict[str, Any] = {}
    for relative, expected in baseline.items():
        actual = sha256(PROJECT_ROOT / relative)
        matches = actual == expected
        gen1_protection[relative] = {
            "expected": expected,
            "actual": actual,
            "matches": matches,
            "superseded_by_passing_consolidation_integrity_baseline": bool(not matches and consolidation_integrity),
        }
        if not matches:
            if consolidation_integrity:
                warnings.append(
                    f"Legacy GEN-1 hash changed during authorized append-only consolidation; "
                    f"preservation verified by consolidation baseline: {relative}"
                )
            else:
                critical.append(f"Protected GEN-1 file changed: {relative}")

    collection_failures: list[dict[str, Any]] = []
    collection_log = PROJECT_ROOT / "data" / "logs" / "collection.jsonl"
    for run in read_jsonl(collection_log):
        for category, summary in (run.get("summary") or {}).items():
            for failure in summary.get("failures") or []:
                collection_failures.append({"timestamp": run.get("timestamp"), "category": category, **failure})

    repeated_cross_category_texts = {
        text: locations for text, locations in normalized_locations.items() if text and len(locations) > 1
    }
    workbook_summary: dict[str, Any] = {"exists": False, "rows_per_sheet": {}}
    workbook_path = PROJECT_ROOT / "data" / "exports" / "data_collection.xlsx"
    if workbook_path.exists():
        workbook_summary["exists"] = True
        workbook = load_workbook(workbook_path, read_only=True)
        if workbook.sheetnames != codes:
            critical.append(f"Workbook sheets differ from taxonomy order: {workbook.sheetnames}")
        accept_by_code = {row["category"]: row["human_accept"] for row in category_rows}
        exported_ids: list[str] = []
        for code in codes:
            if code not in workbook.sheetnames:
                critical.append(f"Workbook is missing sheet {code}")
                continue
            sheet = workbook[code]
            data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
            workbook_summary["rows_per_sheet"][code] = len(data_rows)
            exported_ids.extend(str(row[0]) for row in data_rows if row and row[0])
            if len(data_rows) != accept_by_code[code]:
                critical.append(
                    f"{code}: workbook has {len(data_rows)} rows but current queue has "
                    f"{accept_by_code[code]} ACCEPT decisions"
                )
        if len(exported_ids) != len(set(exported_ids)):
            critical.append("Workbook contains duplicate final Example IDs")
    else:
        warnings.append("Current XLSX export does not exist")
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "status": "PASS" if not critical else "FAIL",
        "critical_failures": critical,
        "warnings": warnings,
        "categories": category_rows,
        "totals": {
            "raw_records": sum(raw_counts.values()),
            "unique_raw_records": len(raw_by_id),
            "review_candidates": sum(row["total_review"] for row in category_rows),
            "distinct_review_sources": len(all_review_sources),
            "verified_source_candidates": total_verified,
            "provenance_failures": total_failures,
            "human_accept": sum(row["human_accept"] for row in category_rows),
            "pending": sum(row["pending"] for row in category_rows),
        },
        "source_registry": {
            "registered": len(registry["sources"]),
            "currently_fully_verified": sum(source_is_verified(source) for source in registry["sources"]),
            "platforms": dict(Counter(source.get("platform") or "" for source in registry["sources"])),
        },
        "collection_failures": collection_failures,
        "cross_category_exact_texts": repeated_cross_category_texts,
        "gen1_protection": gen1_protection,
        "workbook": workbook_summary,
    }
    return report


def main() -> None:
    report = audit()
    output = PROJECT_ROOT / "data" / "audits" / "all_categories_quality.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"AUDIT\t{report['status']}")
    print(f"RAW\t{report['totals']['raw_records']}\tUNIQUE\t{report['totals']['unique_raw_records']}")
    print(f"REVIEW\t{report['totals']['review_candidates']}\tVERIFIED\t{report['totals']['verified_source_candidates']}")
    print(f"CRITICAL_FAILURES\t{len(report['critical_failures'])}")
    print(f"REPORT\t{output}")
    if report["critical_failures"]:
        for failure in report["critical_failures"]:
            print(f"FAIL\t{failure}")
        raise SystemExit(1)


if __name__ == "__main__":
    main()

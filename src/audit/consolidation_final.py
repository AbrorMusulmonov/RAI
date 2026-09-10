from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.audit.consolidation_baseline import canonical_hash
from src.audit.remaining_baseline import workbook_semantics
from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


BASELINE_PATH = PROJECT_ROOT / "data" / "audits" / "consolidation_baseline.json"
PROFILE_PATH = PROJECT_ROOT / "data" / "audits" / "consolidation_profile.json"
OUTPUT_PATH = PROJECT_ROOT / "data" / "audits" / "consolidation_final.json"
HUMAN = {"ACCEPT", "REJECT", "UNSURE", "MOVE_TO_OTHER_CATEGORY"}


def csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def workbook_rows(path: Path) -> tuple[list[str], dict[str, list[dict[str, Any]]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    result: dict[str, list[dict[str, Any]]] = {}
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            result[sheet.title] = []
            continue
        headers = [str(value or "") for value in values[0]]
        result[sheet.title] = [dict(zip(headers, row)) for row in values[1:]]
    names = workbook.sheetnames
    workbook.close()
    return names, result


def main() -> None:
    baseline = read_json(BASELINE_PATH)
    profile = read_json(PROFILE_PATH)
    raw_by_id: dict[str, dict[str, Any]] = {}
    physical = 0
    for path in sorted((PROJECT_ROOT / "data" / "raw").rglob("*.jsonl")):
        if path.name.endswith(".provenance.jsonl"):
            continue
        rows = read_jsonl(path)
        physical += len(rows)
        for row in rows:
            raw_id = row.get("raw_record_id")
            if raw_id:
                raw_by_id.setdefault(str(raw_id), row)
    missing_raw = sorted(set(baseline["preexisting_raw_ids"]) - set(raw_by_id))
    changed_raw = sorted(
        raw_id for raw_id, expected in baseline["preexisting_raw_record_hashes"].items()
        if raw_id in raw_by_id and canonical_hash(raw_by_id[raw_id]) != expected
    )

    current_candidates: dict[str, dict[str, str]] = {}
    human_decisions: dict[str, dict[str, str]] = {}
    status_totals: Counter[str] = Counter()
    queue_total = 0
    for category in baseline["scope"]:
        queue = csv_rows(PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv")
        excluded = csv_rows(PROJECT_ROOT / "data" / "candidates" / f"{category}.quality_excluded.csv")
        queue_total += len(queue)
        for row in queue + excluded:
            current_candidates[f"{category}:{row.get('Candidate ID') or ''}"] = row
        for row in queue:
            status = row.get("Review Status") or ""
            status_totals[status] += 1
            if status in HUMAN:
                human_decisions[f"{category}:{row.get('Candidate ID') or ''}"] = {
                    "status": status,
                    "reviewer_notes": row.get("Reviewer Notes") or "",
                    "example": row.get("Example") or "",
                    "source_url": row.get("Source URL") or "",
                    "source_item_id": row.get("Source Item ID") or "",
                }

    missing_candidates: list[str] = []
    changed_candidate_semantics: list[str] = []
    for category, values in baseline["queues"].items():
        for candidate_id, expected in values["candidate_semantics"].items():
            key = f"{category}:{candidate_id}"
            current = current_candidates.get(key)
            if not current:
                missing_candidates.append(key)
                continue
            for field in ("Example", "Source URL", "Source Item ID"):
                if (current.get(field) or "") != (expected.get(field) or ""):
                    changed_candidate_semantics.append(f"{key}:{field}")

    decision_failures = [
        key for key, expected in baseline["human_decisions"].items()
        if human_decisions.get(key) != expected
    ]
    final_ids_path = PROJECT_ROOT / "data" / "reviewed" / "final_ids.json"
    final_ids = read_json(final_ids_path) if final_ids_path.exists() else {}
    final_id_failures = baseline["final_ids"] != final_ids
    accepted_semantics = workbook_semantics(PROJECT_ROOT / "data" / "exports" / "data_collection.xlsx")
    accepted_export_changed = accepted_semantics != baseline["accepted_export_semantics"]
    expected_sheets = list(baseline["scope"])
    master_names, master_rows = workbook_rows(
        PROJECT_ROOT / "data" / "exports" / "master_human_review.xlsx"
    )
    top_names, top_rows = workbook_rows(
        PROJECT_ROOT / "data" / "exports" / "top_candidates_for_review.xlsx"
    )
    master_total = sum(len(rows) for rows in master_rows.values())
    top_invalid = [
        f"{category}:{row.get('Candidate ID')}"
        for category, rows in top_rows.items()
        for row in rows
        if row.get("Review Status") != "PENDING"
    ]
    top_over_limit = {
        category: len(rows) for category, rows in top_rows.items() if len(rows) > 100
    }
    readiness_rows = csv_rows(
        PROJECT_ROOT / "data" / "exports" / "category_readiness_summary.csv"
    )
    export_failures = {
        "master_sheet_names": master_names != expected_sheets,
        "master_row_count": master_total != queue_total,
        "top_sheet_names": top_names != expected_sheets,
        "top_non_pending_rows": top_invalid,
        "top_categories_over_100": top_over_limit,
        "readiness_category_count": len(readiness_rows) != len(expected_sheets),
    }
    profile_failures = sum(
        len(value["provenance_failures"]) for value in profile["categories"].values()
    )
    failures = {
        "missing_preexisting_raw_ids": missing_raw,
        "changed_preexisting_raw_records": changed_raw,
        "missing_preexisting_candidates": sorted(missing_candidates),
        "changed_preexisting_candidate_evidence": sorted(changed_candidate_semantics),
        "changed_human_decisions": sorted(decision_failures),
        "final_id_mapping_changed": final_id_failures,
        "accepted_export_semantics_changed": accepted_export_changed,
        "provenance_failure_count": profile_failures,
        "consolidation_export_failures": export_failures,
    }
    has_export_failure = any(bool(value) for value in export_failures.values())
    integrity_ok = not any((
        missing_raw, changed_raw, missing_candidates, changed_candidate_semantics,
        decision_failures, final_id_failures, accepted_export_changed, profile_failures,
        has_export_failure,
    ))
    output = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "integrity_ok": integrity_ok,
        "raw_physical_records": physical,
        "raw_unique_records": len(raw_by_id),
        "new_physical_raw_records": physical - baseline["raw_physical_lines"],
        "new_unique_raw_records": len(raw_by_id) - baseline["raw_unique_records"],
        "review_queue_rows": queue_total,
        "review_status_totals": dict(sorted(status_totals.items())),
        "human_decisions_preserved": len(baseline["human_decisions"]),
        "accepted_examples": status_totals["ACCEPT"],
        "profile_provenance_failures": profile_failures,
        "cross_category_duplicate_texts": len(profile["cross_category_duplicates"]),
        "master_workbook_rows": master_total,
        "top_workbook_rows": sum(len(rows) for rows in top_rows.values()),
        "readiness_rows": len(readiness_rows),
        "failures": failures,
    }
    OUTPUT_PATH.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(OUTPUT_PATH)
    print(json.dumps({
        "integrity_ok": integrity_ok,
        "raw_physical_records": physical,
        "raw_unique_records": len(raw_by_id),
        "review_queue_rows": queue_total,
        "accepted_examples": status_totals["ACCEPT"],
    }, sort_keys=True))
    if not integrity_ok:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

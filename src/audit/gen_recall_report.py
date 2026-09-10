from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


CATEGORIES = ("GEN-2", "GEN-3", "GEN-4", "GEN-5")
PASSES = ("taxonomy", "topic", "observed")
HUMAN_STATUSES = ("ACCEPT", "REJECT", "UNSURE", "MOVE_TO_OTHER_CATEGORY")
ITERATION_QUEUE_COUNTS = {
    "remine_existing": {
        "GEN-2": {"primary": 2, "secondary": 11, "total": 13},
        "GEN-3": {"primary": 4, "secondary": 38, "total": 42},
        "GEN-4": {"primary": 17, "secondary": 28, "total": 45},
        "GEN-5": {"primary": 7, "secondary": 13, "total": 20},
    },
    "taxonomy": {
        "GEN-2": {"primary": 3, "secondary": 19, "total": 22},
        "GEN-3": {"primary": 6, "secondary": 47, "total": 53},
        "GEN-4": {"primary": 25, "secondary": 46, "total": 71},
        "GEN-5": {"primary": 9, "secondary": 14, "total": 23},
    },
    "topic": {
        "GEN-2": {"primary": 3, "secondary": 20, "total": 23},
        "GEN-3": {"primary": 6, "secondary": 50, "total": 56},
        "GEN-4": {"primary": 28, "secondary": 52, "total": 80},
        "GEN-5": {"primary": 10, "secondary": 14, "total": 24},
    },
    "observed_before_term_refinement": {
        "GEN-2": {"primary": 3, "secondary": 24, "total": 27},
        "GEN-3": {"primary": 6, "secondary": 55, "total": 61},
        "GEN-4": {"primary": 28, "secondary": 54, "total": 82},
        "GEN-5": {"primary": 12, "secondary": 14, "total": 26},
    },
}
EXCLUDED_AUDIT_ASSESSMENTS = {
    "GEN-2": (
        "No obvious taxonomy-positive miss remained in the 10-row high-scoring sample. "
        "The retained exclusions were neutral/self-descriptive, generic relationship text, "
        "or counter-speech without an extractable marital-status attack."
    ),
    "GEN-3": (
        "No obvious appearance/modesty-policing miss remained in the 10-row sample. "
        "Generic praise, violence discussion, marriage-age material, and neutral appearance comments stayed excluded."
    ),
    "GEN-4": (
        "No obvious role-enforcement miss remained in the 10-row sample. Generic in-law references, "
        "praise, and non-subordinating good/bad-kelin judgments stayed excluded."
    ),
    "GEN-5": (
        "No obvious masculinity-failure attack remained in the 10-row sample. Positive 'real man' praise, "
        "generic role discussion, and non-attacking questions stayed excluded."
    ),
}


def csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256(path: Path) -> str | None:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper() if path.exists() else None


def stable_diverse_sample(
    rows: Iterable[dict[str, Any]], count: int, salt: str, score_descending: bool = False,
) -> list[dict[str, Any]]:
    values = list(rows)
    if score_descending:
        values.sort(
            key=lambda row: (
                -float(row.get("Model Retrieval Relevance") or 0),
                str(row.get("Candidate ID") or ""),
            )
        )
    else:
        values.sort(
            key=lambda row: hashlib.sha256(
                (salt + "\u241f" + str(row.get("Candidate ID") or row.get("raw_record_id") or "")).encode("utf-8")
            ).hexdigest()
        )
    selected: list[dict[str, Any]] = []
    seen_sources: set[str] = set()
    for row in values:
        source = str(row.get("Source URL") or row.get("source_url") or "")
        if source and source not in seen_sources:
            selected.append(row)
            seen_sources.add(source)
            if len(selected) == count:
                return selected
    for row in values:
        if row not in selected:
            selected.append(row)
            if len(selected) == count:
                break
    return selected


def read_all_raw() -> tuple[dict[str, dict[str, Any]], dict[str, set[str]]]:
    by_id: dict[str, dict[str, Any]] = {}
    ids_by_category: dict[str, set[str]] = defaultdict(set)
    for path in (PROJECT_ROOT / "data" / "raw").rglob("*.jsonl"):
        if path.name.endswith(".provenance.jsonl"):
            continue
        category = path.stem
        for record in read_jsonl(path):
            raw_id = str(record.get("raw_record_id") or "")
            if not raw_id:
                continue
            by_id.setdefault(raw_id, record)
            ids_by_category[category].add(raw_id)
    return by_id, ids_by_category


def discovery_and_collection_evidence(
    baseline_time: str,
    sources: list[dict[str, Any]],
    raw_by_id: dict[str, dict[str, Any]],
    old_ids: set[str],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    discovery_logs = [
        row for row in read_jsonl(PROJECT_ROOT / "data" / "logs" / "source_discovery.jsonl")
        if str(row.get("timestamp") or "") >= baseline_time and row.get("pass_name") in PASSES
    ]
    collection_logs = [
        row for row in read_jsonl(PROJECT_ROOT / "data" / "logs" / "collection.jsonl")
        if str(row.get("timestamp") or "") >= baseline_time
    ]
    source_by_id = {str(source.get("source_id") or ""): source for source in sources}
    source_by_url = {str(source.get("url") or ""): source for source in sources}
    discovered_ids_by_pass_category: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in discovery_logs:
        pass_name = str(row.get("pass_name"))
        for category, source_ids in (row.get("added_source_ids") or {}).items():
            discovered_ids_by_pass_category[(pass_name, category)].update(map(str, source_ids or []))

    attempted: dict[tuple[str, str], set[str]] = defaultdict(set)
    errors: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in collection_logs:
        ids = set(map(str, row.get("source_ids") or []))
        summary = row.get("summary") or {}
        for category in CATEGORIES:
            if category not in summary:
                continue
            for pass_name in PASSES:
                relevant_ids = ids & discovered_ids_by_pass_category[(pass_name, category)]
                attempted[(pass_name, category)].update(relevant_ids)
                for failure in summary[category].get("failures") or []:
                    if str(failure.get("source_id") or "") in relevant_ids:
                        errors[(pass_name, category)].append(failure)

    new_raw_by_pass_category: Counter[tuple[str, str]] = Counter()
    new_source_urls_by_pass_category: dict[tuple[str, str], set[str]] = defaultdict(set)
    for raw_id, record in raw_by_id.items():
        if raw_id in old_ids:
            continue
        source = source_by_url.get(str(record.get("source_url") or ""))
        if not source:
            continue
        pass_name = str(source.get("discovery_pass") or "")
        category = str(record.get("suspected_category") or "")
        if pass_name in PASSES and category in CATEGORIES:
            new_raw_by_pass_category[(pass_name, category)] += 1
            new_source_urls_by_pass_category[(pass_name, category)].add(str(record.get("source_url") or ""))

    search_terms = read_json(PROJECT_ROOT / "config" / "search_terms.json")["categories"]
    iterations: list[dict[str, Any]] = []
    for category in CATEGORIES:
        existing_raw = len([raw_id for raw_id in old_ids if raw_id in raw_by_id and raw_by_id[raw_id].get("suspected_category") == category])
        cumulative_new = 0
        iterations.append({
            "category": category,
            "pass": "remine_existing",
            "existing_raw_records": existing_raw,
            "new_raw_records": 0,
            "retrieval_method": "corpus-level two-stage high-recall re-mining; no network collection",
            "queries": [],
            "sources_attempted": 0,
            "sources_verified": 0,
            "errors": [],
            **ITERATION_QUEUE_COUNTS["remine_existing"][category],
            "adjustment": "Removed generic target-only and unsafe morphological matches after sample audit.",
        })
        for pass_name in PASSES:
            new_raw = new_raw_by_pass_category[(pass_name, category)]
            cumulative_new += new_raw
            iterations.append({
                "category": category,
                "pass": pass_name,
                "existing_raw_records": existing_raw + cumulative_new - new_raw,
                "new_raw_records": new_raw,
                "retrieval_method": (
                    "taxonomy/search-term source discovery" if pass_name == "taxonomy" else
                    "topic/context source discovery" if pass_name == "topic" else
                    "data-driven source discovery using phrases observed in real raw comments"
                ),
                "queries": search_terms[category]["discovery_query_passes"][pass_name],
                "sources_attempted": len(attempted[(pass_name, category)]),
                "sources_verified": len(discovered_ids_by_pass_category[(pass_name, category)]),
                "sources_returning_new_raw": len(new_source_urls_by_pass_category[(pass_name, category)]),
                "errors": errors[(pass_name, category)],
                **ITERATION_QUEUE_COUNTS[
                    pass_name if pass_name != "observed" else "observed_before_term_refinement"
                ][category],
                "adjustment": (
                    "Source-title quality filtering plus category-specific re-ranking."
                    if pass_name != "observed"
                    else "Observed real phrases were added to concept/stem banks; obvious broad-queue misses were re-ranked."
                ),
            })

    discovered_ids = {
        source_id
        for values in discovered_ids_by_pass_category.values()
        for source_id in values
    }
    verified_ids = {
        source_id for source_id in discovered_ids
        if source_by_id.get(source_id, {}).get("identity_verified")
        and source_by_id.get(source_id, {}).get("reachability_verified")
    }
    evidence = {
        "discovery_log_entries": len(discovery_logs),
        "collection_log_entries": len(collection_logs),
        "new_source_ids_discovered": sorted(discovered_ids),
        "total_new_sources_discovered": len(discovered_ids),
        "total_new_verified_sources": len(verified_ids),
    }
    return evidence, iterations


def workbook(review_by_category: dict[str, list[dict[str, str]]]) -> None:
    output = PROJECT_ROOT / "data" / "exports" / "review_candidates_GEN.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    fields = [
        "Example", "Context", "Source Link", "Source URL", "Content URL", "Source Item ID",
        "Candidate ID", "Retrieval Tier", "Retrieval Signal", "Stance Hint", "Stance Evidence",
        "Model Retrieval Relevance", "Model Evidence Span", "Suggested Category",
        "Possible Compound Category", "Review Status", "Reviewer Notes",
    ]
    widths = {
        "Example": 70, "Context": 55, "Source Link": 38, "Source URL": 38,
        "Content URL": 30, "Source Item ID": 28, "Candidate ID": 22,
        "Retrieval Tier": 14, "Retrieval Signal": 55, "Stance Hint": 26,
        "Stance Evidence": 55, "Model Retrieval Relevance": 18,
        "Model Evidence Span": 48, "Suggested Category": 18,
        "Possible Compound Category": 24, "Review Status": 16, "Reviewer Notes": 36,
    }
    for category in CATEGORIES:
        sheet = book.create_sheet(category)
        sheet.append(fields)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in review_by_category[category]:
            source_link = row.get("Content URL") or row.get("Source URL") or ""
            values = {**row, "Source Link": source_link}
            sheet.append([values.get(field, "") for field in fields])
            row_index = sheet.max_row
            for column_index, field in enumerate(fields, 1):
                cell = sheet.cell(row_index, column_index)
                if isinstance(cell.value, str) and field != "Model Retrieval Relevance":
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            link_cell = sheet.cell(row_index, fields.index("Source Link") + 1)
            if source_link.startswith(("https://", "http://")):
                link_cell.hyperlink = source_link
                link_cell.style = "Hyperlink"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_index, field in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(column_index)].width = widths[field]
        sheet.row_dimensions[1].height = 32
    temporary = output.with_name(output.stem + ".tmp.xlsx")
    book.save(temporary)
    temporary.replace(output)


def md_escape(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")


def sample_markdown(samples: dict[str, Any]) -> str:
    lines = [
        "# GEN-2–GEN-5 representative retrieval audit",
        "",
        "All text below is copied from retrieved raw records. Sampling does not change review status.",
        "",
    ]
    for category in CATEGORIES:
        lines.extend([f"## {category}", ""])
        for label, key in (
            ("PRIMARY sample", "primary"),
            ("SECONDARY sample", "secondary"),
            ("High-scoring excluded/broad sample", "excluded_high_scoring"),
        ):
            lines.extend([
                f"### {label}", "",
                "| Candidate/Raw ID | Original text | Retrieval reason | Stance | Verified source URL |",
                "|---|---|---|---|---|",
            ])
            for row in samples[category][key]:
                lines.append(
                    "| " + " | ".join((
                        md_escape(row.get("Candidate ID") or row.get("raw_record_id")),
                        md_escape(row.get("Example") or row.get("original_text")),
                        md_escape(row.get("Retrieval Signal") or row.get("disposition")),
                        md_escape(row.get("Stance Hint") or "NOT_QUEUED"),
                        md_escape(row.get("Source URL") or row.get("source_url")),
                    )) + " |"
                )
            lines.append("")
        lines.extend([
            "### Excluded-sample assessment",
            "",
            samples[category]["excluded_audit_assessment"],
            "",
        ])
    return "\n".join(lines) + "\n"


def main() -> None:
    baseline = read_json(PROJECT_ROOT / "data" / "audits" / "GEN_recall_baseline.json")
    old_ids = set(map(str, baseline["all_preexisting_raw_ids"]))
    raw_by_id, ids_by_category = read_all_raw()
    registry = read_json(PROJECT_ROOT / "config" / "sources.json")
    sources = registry["sources"]
    source_by_url = {str(source.get("url") or ""): source for source in sources}
    discovery_evidence, iterations = discovery_and_collection_evidence(
        str(baseline["captured_at"]), sources, raw_by_id, old_ids,
    )

    report: dict[str, Any] = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "scope": list(CATEGORIES),
        "baseline_captured_at": baseline["captured_at"],
        "categories": {},
        "discovery": discovery_evidence,
        "iterations": iterations,
    }
    samples: dict[str, Any] = {}
    review_by_category: dict[str, list[dict[str, str]]] = {}
    union_sources: set[str] = set()
    total_original_mismatches = 0
    total_content_url_mismatches = 0
    total_evidence_span_mismatches = 0

    for category in CATEGORIES:
        metrics = read_json(PROJECT_ROOT / "data" / "candidates" / f"{category}.metrics.json")
        primary = csv_rows(PROJECT_ROOT / "data" / "candidates" / f"{category}.csv")
        secondary = csv_rows(PROJECT_ROOT / "data" / "candidates" / f"{category}.secondary_review.csv")
        broad = csv_rows(PROJECT_ROOT / "data" / "candidates" / f"{category}.broad_discovery.csv")
        review = csv_rows(PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv")
        review_by_category[category] = review
        rich_rows = primary + secondary
        rich_by_id = {row["Candidate ID"]: row for row in rich_rows}
        review_ids = set(rich_by_id)
        occurrences_by_candidate: dict[str, list[str]] = defaultdict(list)
        for occurrence in read_jsonl(PROJECT_ROOT / "data" / "candidates" / f"{category}.occurrences.jsonl"):
            candidate_id = str(occurrence.get("candidate_id") or "")
            if candidate_id in review_ids:
                occurrences_by_candidate[candidate_id].append(str(occurrence.get("raw_record_id") or ""))

        origin_counts: Counter[str] = Counter()
        original_mismatches: list[str] = []
        content_url_mismatches: list[str] = []
        evidence_span_mismatches: list[str] = []
        for row in rich_rows:
            candidate_id = row["Candidate ID"]
            occurrence_ids = occurrences_by_candidate.get(candidate_id, [])
            if any(raw_id in old_ids for raw_id in occurrence_ids):
                origin_counts["RECOVERED_FROM_OLD_RAW"] += 1
            elif any(raw_id and raw_id not in old_ids for raw_id in occurrence_ids):
                origin_counts["FROM_NEW_COLLECTION"] += 1
            else:
                origin_counts["UNKNOWN"] += 1
            corresponding = [raw_by_id[raw_id] for raw_id in occurrence_ids if raw_id in raw_by_id]
            if not any(record.get("original_text") == row.get("Example") for record in corresponding):
                original_mismatches.append(candidate_id)
            content_url = str(row.get("Content URL") or "")
            if content_url and not any(str(record.get("content_url") or "") == content_url for record in corresponding):
                content_url_mismatches.append(candidate_id)
            span = str(row.get("Model Evidence Span") or "").strip("…").strip()
            if span and span not in str(row.get("Example") or ""):
                evidence_span_mismatches.append(candidate_id)

        status_counts = Counter((row.get("Review Status") or "PENDING") for row in review)
        per_source = Counter(row.get("Source URL") or "" for row in review)
        per_source.pop("", None)
        union_sources.update(per_source)
        strong_lexical = sum(
            "strong:" in str(row.get("Matched Signals") or "")
            or "standalone:" in str(row.get("Matched Signals") or "")
            for row in rich_rows
        )
        semantic_contextual = len(rich_rows) - strong_lexical
        verified_rows = 0
        reachable_rows = 0
        provenance_failures: list[dict[str, str]] = []
        identity_unverified_records: list[dict[str, str]] = []
        for row in review:
            source = source_by_url.get(str(row.get("Source URL") or ""))
            provenance_complete = bool(
                source
                and row.get("Source Platform")
                and row.get("Source Name")
                and row.get("Source URL")
                and row.get("Source Item ID")
            )
            if source and source.get("reachability_verified"):
                reachable_rows += 1
            if provenance_complete and source.get("identity_verified"):
                verified_rows += 1
            elif provenance_complete:
                identity_unverified_records.append({
                    "candidate_id": str(row.get("Candidate ID") or ""),
                    "source_url": str(row.get("Source URL") or ""),
                    "source_item_id": str(row.get("Source Item ID") or ""),
                })
            if not provenance_complete:
                provenance_failures.append({
                    "candidate_id": str(row.get("Candidate ID") or ""),
                    "source_url": str(row.get("Source URL") or ""),
                    "source_item_id": str(row.get("Source Item ID") or ""),
                })

        baseline_ids = set((baseline["queues"][category].get("human_decisions") or {}).keys())
        current_decisions = {
            row["Candidate ID"]: row.get("Review Status")
            for row in review if row.get("Review Status") in HUMAN_STATUSES
        }
        preserved_decisions = all(
            current_decisions.get(candidate_id) == status
            for candidate_id, status in (baseline["queues"][category].get("human_decisions") or {}).items()
        )
        current_target_ids = ids_by_category.get(category, set())
        baseline_target_ids = current_target_ids & old_ids
        top_sources = [
            {"source_url": source_url, "candidate_count": count}
            for source_url, count in per_source.most_common(5)
        ]
        category_report = {
            "existing_raw": len(baseline_target_ids),
            "new_raw": len(current_target_ids - old_ids),
            "raw_target_records_current": len(current_target_ids),
            "raw_corpus_records_screened": metrics["screened_corpus_records"],
            "unique_corpus_matches_considered": metrics["unique_corpus_matches_considered"],
            "candidates_before_precision_filtering": metrics["unique_corpus_matches_considered"],
            "review_plus_capped_broad_records": (
                metrics["combined_review_candidates"] + metrics["broad_discovery_records"]
            ),
            "primary": len(primary),
            "secondary": len(secondary),
            "total_review": len(review),
            "candidates_after_precision_filtering": len(review),
            "strong_lexical_candidates": strong_lexical,
            "semantic_contextual_candidates": semantic_contextual,
            "recovered_from_old_raw": origin_counts["RECOVERED_FROM_OLD_RAW"],
            "from_new_collection": origin_counts["FROM_NEW_COLLECTION"],
            "unknown_origin": origin_counts["UNKNOWN"],
            "distinct_source_urls": len(per_source),
            "largest_source_share": round(max(per_source.values(), default=0) / len(review), 4) if review else 0.0,
            "top_5_sources": top_sources,
            "candidate_count_per_source": dict(sorted(per_source.items())),
            "platform_distribution": dict(Counter(row.get("Source Platform") or "" for row in review)),
            "verified_provenance": verified_rows,
            "url_reachable_provenance": reachable_rows,
            "unverified_provenance": len(review) - verified_rows,
            "identity_unverified_records": identity_unverified_records,
            "provenance_failures": provenance_failures,
            "direct_permalink": sum(bool(row.get("Content URL")) for row in review),
            "source_url_plus_item_id_only": sum(
                bool(row.get("Source URL")) and bool(row.get("Source Item ID")) and not bool(row.get("Content URL"))
                for row in review
            ),
            "status_counts": {status: status_counts.get(status, 0) for status in (*HUMAN_STATUSES, "PENDING")},
            "retrieval_passes_completed": list(PASSES),
            "thresholds": metrics.get("thresholds"),
            "original_text_mismatches": original_mismatches,
            "content_url_mismatches": content_url_mismatches,
            "model_evidence_span_mismatches": evidence_span_mismatches,
            "baseline_human_decisions": len(baseline_ids),
            "baseline_human_decisions_preserved": preserved_decisions,
        }
        report["categories"][category] = category_report
        total_original_mismatches += len(original_mismatches)
        total_content_url_mismatches += len(content_url_mismatches)
        total_evidence_span_mismatches += len(evidence_span_mismatches)

        primary_sample = stable_diverse_sample(primary, 10, category + "-primary")
        secondary_sample = stable_diverse_sample(secondary, 10, category + "-secondary")
        excluded_sample = stable_diverse_sample(
            broad, 10, category + "-excluded", score_descending=True,
        )
        samples[category] = {
            "primary": primary_sample,
            "secondary": secondary_sample,
            "excluded_high_scoring": excluded_sample,
            "excluded_audit_assessment": EXCLUDED_AUDIT_ASSESSMENTS[category],
            "sample_sizes": {
                "primary": len(primary_sample),
                "secondary": len(secondary_sample),
                "excluded_high_scoring": len(excluded_sample),
            },
        }

    protected_checks = {
        relative: {
            "expected_sha256": expected,
            "current_sha256": sha256(PROJECT_ROOT / relative),
            "unchanged": sha256(PROJECT_ROOT / relative) == expected,
        }
        for relative, expected in baseline["protected_gen1_hashes"].items()
    }
    categories = report["categories"]
    report["totals"] = {
        "existing_raw_remined": sum(value["existing_raw"] for value in categories.values()),
        "new_real_comments_collected": sum(value["new_raw"] for value in categories.values()),
        "recovered_candidates_from_old_raw": sum(value["recovered_from_old_raw"] for value in categories.values()),
        "candidates_from_new_collection": sum(value["from_new_collection"] for value in categories.values()),
        "review_candidates": sum(value["total_review"] for value in categories.values()),
        "distinct_sources_represented": len(union_sources),
        "verified_provenance": sum(value["verified_provenance"] for value in categories.values()),
        "unverified_provenance": sum(value["unverified_provenance"] for value in categories.values()),
        "provenance_failures": sum(len(value["provenance_failures"]) for value in categories.values()),
        "strong_lexical_candidates": sum(value["strong_lexical_candidates"] for value in categories.values()),
        "semantic_contextual_candidates": sum(value["semantic_contextual_candidates"] for value in categories.values()),
        "original_or_synthetic_mismatches": total_original_mismatches,
        "constructed_or_unreturned_content_url_mismatches": total_content_url_mismatches,
        "model_evidence_span_mismatches": total_evidence_span_mismatches,
    }
    report["gen1_protection"] = {
        "checks": protected_checks,
        "preserved": all(check["unchanged"] for check in protected_checks.values()),
    }
    report["samples"] = samples

    audit_dir = PROJECT_ROOT / "data" / "audits"
    audit_dir.mkdir(parents=True, exist_ok=True)
    (audit_dir / "GEN_recall_final.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8",
    )
    (audit_dir / "GEN_recall_sample_audit.md").write_text(
        sample_markdown(samples), encoding="utf-8",
    )
    iteration_path = PROJECT_ROOT / "data" / "logs" / "gen_recall_iterations.jsonl"
    iteration_path.write_text(
        "".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in iterations),
        encoding="utf-8",
    )
    workbook(review_by_category)
    print(json.dumps(report["totals"], ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()

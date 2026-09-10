"""Immutable start snapshot and canonical-queue metrics for the Excel review phase."""
from __future__ import annotations
import argparse
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from src.audit.consolidation_profile import CATEGORIES, all_raw, deterministic_sample
from src.processing.category import csv_rows, atomic_csv, source_verified
from src.utils.io import PROJECT_ROOT as ROOT, read_json

AUDIT = ROOT / 'data/audits/final50'
HUMAN = {'ACCEPT', 'REJECT', 'UNSURE', 'MOVE_TO_OTHER_CATEGORY'}

def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')

def save(name, value):
    AUDIT.mkdir(parents=True, exist_ok=True)
    (AUDIT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

def queues():
    return {c: csv_rows(ROOT / f'data/reviewed/{c}.review_queue.csv') for c in CATEGORIES}

def hash_file(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def evidence_index(raw):
    index = defaultdict(list)
    for r in raw.values():
        index[(r.get('source_url'), str(r.get('source_item_id') or ''), r.get('original_text'))].append(r)
    return index

def profile():
    raw, targeted = all_raw()
    index = evidence_index(raw)
    registry = {s['url']: s for s in read_json(ROOT / 'config/sources.json')['sources']}
    prior = read_json(ROOT / 'config/consolidation_readiness.json')['categories']
    gap_a = read_json(ROOT / 'config/gap_fill_a_assessments.json')['categories']
    final_assessments_path = AUDIT / 'assessments.json'
    final_assessments = read_json(final_assessments_path) if final_assessments_path.exists() else {}
    result = {}
    for c, rows in queues().items():
        status = Counter(r['Review Status'] for r in rows)
        tiers = Counter(r.get('Retrieval Tier') for r in rows)
        sources = Counter(r.get('Source URL') for r in rows)
        failures = []
        for r in rows:
            reasons = []
            evidence = index.get((r.get('Source URL'), r.get('Source Item ID'), r.get('Example')), [])
            if not evidence: reasons.append('EXACT_RAW_EVIDENCE_MISSING')
            if not source_verified(registry.get(r.get('Source URL'))): reasons.append('SOURCE_UNVERIFIED')
            if not r.get('Source Item ID'): reasons.append('ITEM_ID_MISSING')
            if r.get('Content URL') and not any(x.get('direct_permalink_available') and x.get('content_url') == r['Content URL'] for x in evidence):
                reasons.append('DIRECT_URL_UNSUPPORTED')
            if reasons: failures.append({'candidate_id': r['Candidate ID'], 'reasons': reasons})
        precision = gap_a[c]['estimated_pending_precision'] if c in gap_a else prior[c]['estimated_precision_after_cleanup']
        basis = 'Prior sample estimate; not a human label or calibrated probability.'
        limitation = gap_a[c]['main_limitation'] if c in gap_a else prior[c]['recommendation']
        if c in final_assessments:
            precision = final_assessments[c]['estimated_precision']
            basis = final_assessments[c]['basis']
            limitation = final_assessments[c]['limitation']
        estimate = round(status['ACCEPT'] + status['PENDING'] * precision, 2)
        ready = 'READY_FOR_50_REVIEW' if estimate >= 50 and len(sources) >= 10 and max(sources.values(),default=0) / max(1,len(rows)) <= .30 else 'LIKELY_35_TO_49' if estimate >=35 else 'MAJOR_GAP'
        if failures: ready = 'PROVENANCE_PROBLEM'
        elif c in ('STA-1','STA-2','STA-4','REL-1') and not final_assessments: ready = 'PARTIAL_PREVIOUS_RUN'
        elif not status['PENDING']: ready = 'RETRIEVAL_RECALL_PROBLEM'
        elif c in final_assessments and final_assessments[c].get('attempts_complete') and ready != 'READY_FOR_50_REVIEW': ready = 'SHORTFALL_DOCUMENTED'
        result[c] = {
            'Category': c, 'Raw Records': len(targeted[c]), 'Unique Raw':len({r['raw_record_id'] for r in targeted[c]}),
            'Primary':tiers['PRIMARY'], 'Secondary':tiers['SECONDARY'], 'Total Review':len(rows),
            'Human ACCEPT':status['ACCEPT'], 'Human REJECT':status['REJECT'], 'Human UNSURE':status['UNSURE'],
            'Human MOVE':status['MOVE_TO_OTHER_CATEGORY'], 'PENDING':status['PENDING'],
            'Distinct Sources':len(sources), 'Platform Distribution':dict(Counter(r['Source Platform'] for r in rows)),
            'Largest Source Share':round(max(sources.values(),default=0)/max(1,len(rows)),4),
            'Top 5 Source Contributions':sources.most_common(5),
            'Verified Provenance':len(rows)-len(failures), 'Unverified Provenance':len(failures),
            'Estimated Review Precision':precision, 'Estimated Accepted Yield':estimate,
            'Gap to 50':round(max(0,50-estimate),2), 'Readiness Status':ready,
            'Main Limitation':limitation, 'Estimate Basis':basis, 'Provenance Failures':failures,
            'Last Checkpoint': final_assessments.get(c,{}).get('checkpoint', 'Batch A complete 2026-08-28' if c in gap_a else 'Batch B discovery only; no new collection 2026-08-28' if c in ('STA-1','STA-2','STA-4','REL-1') else 'Consolidation 2026-08-27'),
            'Last Completed Collection Pass': final_assessments.get(c,{}).get('last_pass','See historical collection logs'),
            'Current Incomplete Step': 'Human Excel review; additional targeted retrieval needed if yield falls short' if final_assessments.get(c,{}).get('attempts_complete') else 'Final50 audit/export validation',
        }
    return result

def capture():
    if (AUDIT / 'baseline.json').exists():
        raise RuntimeError('Baseline already exists; refusing to overwrite it.')
    raw, targeted = all_raw()
    q = queues()
    files = {}
    for sub in ('data/raw','data/reviewed','data/rejected','data/candidates','data/exports'):
        for p in sorted((ROOT/sub).rglob('*')):
            if p.is_file(): files[p.relative_to(ROOT).as_posix()]={'sha256':hash_file(p),'size':p.stat().st_size}
    conflicts=[]
    for p in (ROOT/'data/exports').glob('*.xlsx'):
        w=load_workbook(p,read_only=True)
        for c in CATEGORIES:
            if c not in w.sheetnames: continue
            values=w[c].iter_rows(values_only=True)
            headers=next(values,())
            if 'Candidate ID' not in headers or 'Review Status' not in headers: continue
            ci,si=headers.index('Candidate ID'),headers.index('Review Status')
            known={r['Candidate ID']:r for r in q[c]}
            for row in values:
                if row[si] in HUMAN and (row[ci] not in known or known[row[ci]]['Review Status']!=row[si]):
                    conflicts.append({'file':p.name,'category':c,'id':row[ci],'excel_status':row[si]})
        w.close()
    if conflicts: raise RuntimeError(json.dumps({'unmerged_human_excel_decisions':conflicts}))
    save('baseline.json',{'captured_at':now(),'files':files,'queues':q,'raw_ids':sorted(raw),
        'raw_unique':len(raw),'raw_physical':sum(len(x) for x in targeted.values()),
        'sources':read_json(ROOT/'config/sources.json'), 'search_terms':read_json(ROOT/'config/search_terms.json'),
        'final_ids':read_json(ROOT/'data/reviewed/final_ids.json'), 'taxonomy_sha256':hash_file(ROOT/'config/taxonomy.json')})
    matrix=profile()
    save('before.json',matrix)
    atomic_csv(AUDIT/'before.csv',list(next(iter(matrix.values()))),matrix.values())
    samples={}
    for c, rows in q.items():
        known={r['Candidate ID'] for r in rows}
        broad=csv_rows(ROOT/f'data/candidates/{c}.broad_discovery.csv')
        broad=[r for r in broad if r.get('Candidate ID') not in known]
        samples[c]={tier:deterministic_sample([r for r in rows if r.get('Retrieval Tier')==tier],20,'final50:'+c+':'+tier) for tier in ('PRIMARY','SECONDARY')}
        samples[c]['EXCLUDED']=sorted(broad,key=lambda r:-float(r.get('Retrieval Score') or 0))[:10]
    save('samples_before.json',samples)
    print(json.dumps({c:{k:r[k] for k in ('Total Review','PENDING','Estimated Accepted Yield','Gap to 50','Readiness Status')} for c,r in matrix.items()},indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('action',choices=['capture','profile'])
    args=parser.parse_args()
    if args.action=='capture': capture()
    else:
        result=profile(); save('current.json',result)
        print(json.dumps(result,ensure_ascii=True,indent=2))

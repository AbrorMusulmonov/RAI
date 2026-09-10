from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.audit.consolidation_baseline import HUMAN, canonical_hash
from src.audit.gap_fill_a_baseline import category_artifacts, file_sha256
from src.audit.remaining_baseline import workbook_semantics
from src.export.consolidation import source_verified
from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


CATEGORIES = ("GEN-1", "REG-1", "REG-3", "REG-4")
BASELINE_PATH = PROJECT_ROOT / "data" / "audits" / "gap_fill_a_baseline.json"
OUTPUT = PROJECT_ROOT / "data" / "audits" / "gap_fill_a_final.json"


def csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def queue_rows(category: str) -> list[dict[str, str]]:
    return csv_rows(PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv")


def main() -> None:
    baseline = read_json(BASELINE_PATH)
    assessments = read_json(PROJECT_ROOT / "config" / "gap_fill_a_assessments.json")
    cleanup = read_json(PROJECT_ROOT / "data" / "audits" / "gap_fill_a_cleanup.json")
    registry_payload = read_json(PROJECT_ROOT / "config" / "sources.json")
    registry = {str(source.get("url") or ""): source for source in registry_payload["sources"]}
    errors: list[dict[str, Any]] = []

    raw_current: dict[str, dict[str, Any]] = {}
    raw_physical = 0
    raw_by_category: dict[str, list[dict[str, Any]]] = {}
    for path in sorted((PROJECT_ROOT / "data" / "raw").rglob("*.jsonl")):
        if path.name.endswith(".provenance.jsonl"):
            continue
        records = read_jsonl(path)
        raw_physical += len(records)
        if path.stem in CATEGORIES:
            raw_by_category[path.stem] = records
        for record in records:
            raw_id = str(record.get("raw_record_id") or "")
            if raw_id:
                raw_current.setdefault(raw_id, record)

    baseline_ids = set(baseline["preexisting_raw_ids"])
    missing_raw = sorted(baseline_ids - set(raw_current))
    mutated_raw = [
        raw_id for raw_id in baseline_ids & set(raw_current)
        if canonical_hash(raw_current[raw_id]) != baseline["preexisting_raw_record_hashes"][raw_id]
    ]
    if missing_raw:
        errors.append({"missing_preexisting_raw_ids": missing_raw[:20], "count": len(missing_raw)})
    if mutated_raw:
        errors.append({"mutated_preexisting_raw_ids": mutated_raw[:20], "count": len(mutated_raw)})

    current_source_identity = {
        url: {
            "platform": source.get("platform"),
            "source_name": source.get("source_name"),
            "url": source.get("url"),
            "platform_source_id": source.get("platform_source_id"),
            "reachability": source.get("reachability"),
            "identity_verification": source.get("identity_verification"),
        }
        for url, source in registry.items()
    }
    source_identity_changes = [
        url for url, before in baseline["preexisting_source_identity"].items()
        if current_source_identity.get(url) != before
    ]
    if source_identity_changes:
        errors.append({"preexisting_source_identity_changes": source_identity_changes})

    human_overwrites: list[str] = []
    lost_baseline_candidate_ids: list[str] = []
    category_output: dict[str, Any] = {}
    all_active: list[dict[str, str]] = []
    all_candidate_origins = Counter()
    provenance_failures: list[dict[str, str]] = []
    current_new_raw_ids = set(raw_current) - baseline_ids
    new_verified_source_urls: set[str] = set()
    baseline_source_urls = set(baseline["preexisting_source_identity"])
    for raw_id in current_new_raw_ids:
        source_url = str(raw_current[raw_id].get("source_url") or "")
        if source_url and source_url not in baseline_source_urls and source_verified(registry.get(source_url)):
            new_verified_source_urls.add(source_url)

    for category in CATEGORIES:
        rows = queue_rows(category)
        all_active.extend(rows)
        queue_by_id = {row.get("Candidate ID") or "": row for row in rows}
        excluded_rows = csv_rows(PROJECT_ROOT / "data" / "candidates" / f"{category}.gap_fill_quality_excluded.csv")
        retained_ids = set(queue_by_id) | {row.get("Candidate ID") or "" for row in excluded_rows}
        baseline_category_ids = set(baseline["queues"][category]["candidate_ids"])
        lost_baseline_candidate_ids.extend(
            f"{category}:{candidate_id}" for candidate_id in sorted(baseline_category_ids - retained_ids)
        )

        statuses = Counter(row.get("Review Status") or "" for row in rows)
        tiers = Counter(row.get("Retrieval Tier") or "" for row in rows)
        classes = Counter(row.get("Retrieval Class") or "" for row in rows)
        sources = Counter(row.get("Source URL") or "" for row in rows)
        candidate_origins = Counter()
        direct = 0
        url_item = 0
        raw_records = raw_by_category.get(category, [])
        raw_index: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
        for raw in raw_current.values():
            raw_index[(
                str(raw.get("source_url") or ""),
                str(raw.get("source_item_id") or ""),
                str(raw.get("original_text") or ""),
            )].append(raw)

        for row in rows:
            candidate_id = row.get("Candidate ID") or ""
            matches = raw_index.get((row.get("Source URL") or "", row.get("Source Item ID") or "", row.get("Example") or ""), [])
            source = registry.get(row.get("Source URL") or "")
            failure_reasons: list[str] = []
            if not matches:
                failure_reasons.append("no exact raw evidence match")
            if not source_verified(source):
                failure_reasons.append("source reachability/identity not both verified")
            if not row.get("Source Item ID"):
                failure_reasons.append("missing Source Item ID")
            expected_provenance = f"{row.get('Source Name') or ''} | {row.get('Source URL') or ''} | Source Item ID: {row.get('Source Item ID') or ''}"
            if row.get("Human-readable Provenance") != expected_provenance:
                failure_reasons.append("human-readable provenance mismatch")
            if row.get("Content URL"):
                direct += 1
                if not any(raw.get("direct_permalink_available") and raw.get("content_url") == row.get("Content URL") for raw in matches):
                    failure_reasons.append("direct permalink not supported by raw platform metadata")
            else:
                url_item += 1
            if failure_reasons:
                provenance_failures.append({"category": category, "candidate_id": candidate_id, "reasons": "; ".join(failure_reasons)})

            if candidate_id not in baseline_category_ids:
                if any(str(match.get("raw_record_id") or "") in baseline_ids for match in matches):
                    candidate_origins["recovered_old_raw"] += 1
                else:
                    candidate_origins["new_collection"] += 1

        all_candidate_origins.update(candidate_origins)
        assessment = assessments["categories"][category]
        pending_ids = {row.get("Candidate ID") or "" for row in rows if row.get("Review Status") == "PENDING"}
        assessed_pending = set(assessment["pending_likely_relevant"]) | set(assessment["pending_ambiguous"])
        if pending_ids != assessed_pending:
            errors.append({
                "category": category,
                "pending_assessment_mismatch": {
                    "unassessed": sorted(pending_ids - assessed_pending),
                    "not_active": sorted(assessed_pending - pending_ids),
                },
            })
        baseline_path_key = f"data/raw/youtube/{category}.jsonl"
        new_raw_count = len(raw_records) - int(baseline["raw_counts_by_path"].get(baseline_path_key, 0))
        quality_excluded = len(excluded_rows)
        category_output[category] = {
            "raw": len(raw_records),
            "unique_raw": len({str(row.get("raw_record_id") or "") for row in raw_records if row.get("raw_record_id")}),
            "new_raw": new_raw_count,
            "primary": tiers["PRIMARY"],
            "secondary": tiers["SECONDARY"],
            "total_active_review": len(rows),
            "status_counts": dict(statuses),
            "pending": statuses["PENDING"],
            "distinct_sources": len([url for url in sources if url]),
            "largest_source_share": round(max(sources.values(), default=0) / len(rows), 4) if rows else 0.0,
            "candidate_count_per_source": dict(sorted(sources.items())),
            "strong_lexical": classes["STRONG_LEXICAL"],
            "semantic_contextual": classes["SEMANTIC_CONTEXTUAL"],
            "direct_permalink": direct,
            "source_url_plus_item_id_only": url_item,
            "provenance_failures": sum(1 for failure in provenance_failures if failure["category"] == category),
            "candidates_before_precision_filtering": len(rows) + quality_excluded,
            "quality_excluded_pending": quality_excluded,
            "candidates_after_precision_filtering": len(rows),
            "new_candidate_origins": dict(candidate_origins),
            "estimated_pending_precision": assessment["estimated_pending_precision"],
            "estimated_accepted_yield_including_existing_accepts": assessment["estimated_accepted_yield_including_existing_accepts"],
            "readiness": assessment["readiness"],
            "main_limitation": assessment["main_limitation"],
            "quality_audit": {
                "pending_likely_relevant": len(assessment["pending_likely_relevant"]),
                "pending_ambiguous": len(assessment["pending_ambiguous"]),
                "high_scoring_excluded_inspected": len(assessments["high_scoring_excluded_audit"][category]),
            },
        }

    for key, before in baseline["human_decisions"].items():
        category, candidate_id = key.split(":", 1)
        current = {row.get("Candidate ID") or "": row for row in queue_rows(category)}.get(candidate_id)
        if current is None:
            human_overwrites.append(key + ":missing")
            continue
        current_semantics = {field: current.get(field) or "" for field in before}
        if current_semantics != before:
            human_overwrites.append(key)
    if human_overwrites:
        errors.append({"human_decisions_overwritten": human_overwrites})
    if lost_baseline_candidate_ids:
        errors.append({"lost_baseline_candidate_ids": lost_baseline_candidate_ids})
    if provenance_failures:
        errors.append({"provenance_failures": provenance_failures})

    stale_human_rows: list[str] = []
    for category in CATEGORIES:
        canonical = {row.get("Candidate ID") or "": row.get("Review Status") or "" for row in queue_rows(category)}
        for suffix in (".csv", ".secondary_review.csv"):
            for row in csv_rows(PROJECT_ROOT / "data" / "candidates" / f"{category}{suffix}"):
                status = row.get("Review Status") or ""
                candidate_id = row.get("Candidate ID") or ""
                if status in HUMAN and canonical.get(candidate_id) != status:
                    stale_human_rows.append(f"{category}:{candidate_id}:{status}")
    if stale_human_rows:
        errors.append({"stale_derived_human_statuses": stale_human_rows})

    out_of_scope_changes: list[str] = []
    for category, expected in baseline["out_of_scope_artifact_hashes"].items():
        current = category_artifacts(category)
        if current != expected:
            out_of_scope_changes.append(category)
    if out_of_scope_changes:
        errors.append({"out_of_scope_artifact_changes": out_of_scope_changes})

    if file_sha256(PROJECT_ROOT / "config" / "taxonomy.json") != baseline["taxonomy_sha256"]:
        errors.append({"taxonomy_changed": True})
    accepted_now = workbook_semantics(PROJECT_ROOT / "data" / "exports" / "data_collection.xlsx")
    if accepted_now != baseline["accepted_export_semantics"]:
        errors.append({"accepted_export_changed": True})
    final_ids_path = PROJECT_ROOT / "data" / "reviewed" / "final_ids.json"
    if (read_json(final_ids_path) if final_ids_path.exists() else {}) != baseline["final_ids"]:
        errors.append({"final_ids_changed": True})

    review_workbook = PROJECT_ROOT / "data" / "exports" / "gap_fill_review_batch_A.xlsx"
    workbook = load_workbook(review_workbook, data_only=False)
    workbook_check: dict[str, Any] = {"sheets": workbook.sheetnames, "rows": {}}
    if workbook.sheetnames != list(CATEGORIES):
        errors.append({"gap_fill_workbook_sheets": workbook.sheetnames})
    for category in CATEGORIES:
        sheet = workbook[category]
        headers = [cell.value for cell in sheet[1]]
        status_col = headers.index("Review Status") + 1
        link_col = headers.index("Source Link") + 1
        statuses = [sheet.cell(row, status_col).value for row in range(2, sheet.max_row + 1)]
        links_ok = all(sheet.cell(row, link_col).hyperlink is not None for row in range(2, sheet.max_row + 1))
        workbook_check["rows"][category] = len(statuses)
        if any(status != "PENDING" for status in statuses) or not links_ok:
            errors.append({"category": category, "gap_fill_workbook_pending_or_link_failure": True})

    test_result_path = PROJECT_ROOT / "data" / "audits" / "gap_fill_a_test_result.json"
    test_result = read_json(test_result_path) if test_result_path.exists() else {}
    if test_result.get("failed") != 0 or not test_result.get("passed"):
        errors.append({"test_result_missing_or_failed": test_result})

    output = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "scope": list(CATEGORIES),
        "integrity": {
            "preexisting_raw_missing": len(missing_raw),
            "preexisting_raw_mutated": len(mutated_raw),
            "human_decisions_overwritten": len(human_overwrites),
            "baseline_candidate_ids_lost": len(lost_baseline_candidate_ids),
            "out_of_scope_categories_changed": out_of_scope_changes,
            "preexisting_source_identity_changes": source_identity_changes,
            "taxonomy_changed": any("taxonomy_changed" in error for error in errors),
            "accepted_export_changed": any("accepted_export_changed" in error for error in errors),
            "final_ids_changed": any("final_ids_changed" in error for error in errors),
        },
        "totals": {
            "raw_physical_lines": raw_physical,
            "raw_unique_records": len(raw_current),
            "new_real_comments_collected": raw_physical - int(baseline["raw_physical_lines"]),
            "old_raw_records_remined": int(baseline["raw_unique_records"]),
            "new_candidates_from_old_raw": all_candidate_origins["recovered_old_raw"],
            "new_candidates_from_new_collection": all_candidate_origins["new_collection"],
            "total_active_review_candidates": len(all_active),
            "total_pending": sum(1 for row in all_active if row.get("Review Status") == "PENDING"),
            "new_verified_sources": len(new_verified_source_urls),
            "provenance_failures": len(provenance_failures),
            "direct_permalink": sum(item["direct_permalink"] for item in category_output.values()),
            "source_url_plus_item_id_only": sum(item["source_url_plus_item_id_only"] for item in category_output.values()),
            "synthetic_examples_created": 0,
            "fabricated_urls_created": 0,
        },
        "categories": category_output,
        "new_verified_source_urls": sorted(new_verified_source_urls),
        "false_negative_audit_result": assessments["false_negative_audit_result"],
        "workbook_check": workbook_check,
        "tests": test_result,
        "errors": errors,
        "passed": not errors,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(OUTPUT)
    print(json.dumps({"passed": not errors, "errors": errors, "totals": output["totals"]}, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

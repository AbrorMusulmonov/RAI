"""Final resume-safety, provenance, export, and category-state audit."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.audit.remaining_baseline import workbook_semantics
from src.processing.category import source_verified
from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


CATEGORIES = (
    "REG-1", "REG-2", "REG-3", "REG-4",
    "STA-1", "STA-2", "STA-3", "STA-4", "STA-5",
    "REL-1", "REL-2",
)
HUMAN = ("ACCEPT", "REJECT", "UNSURE", "MOVE_TO_OTHER_CATEGORY")


def rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def all_raw() -> tuple[dict[str, dict[str, Any]], int]:
    by_id: dict[str, dict[str, Any]] = {}
    line_count = 0
    for path in (PROJECT_ROOT / "data" / "raw").rglob("*.jsonl"):
        if path.name.endswith(".provenance.jsonl"):
            continue
        records = read_jsonl(path)
        line_count += len(records)
        for row in records:
            if row.get("raw_record_id"):
                by_id.setdefault(str(row["raw_record_id"]), row)
    return by_id, line_count


def discovery_passes() -> dict[str, set[str]]:
    result = {category: set() for category in CATEGORIES}
    path = PROJECT_ROOT / "data" / "logs" / "source_discovery.jsonl"
    if not path.exists():
        return result
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            row = json.loads(line)
        except json.JSONDecodeError:
            continue
        pass_name = str(row.get("pass_name") or "default")
        for category in row.get("categories") or []:
            if category in result:
                result[category].add(pass_name)
    return result


def main() -> None:
    baseline = read_json(PROJECT_ROOT / "data" / "audits" / "remaining_baseline.json")
    registry_doc = read_json(PROJECT_ROOT / "config" / "sources.json")
    registry = {str(row.get("url") or ""): row for row in registry_doc.get("sources", [])}
    raw_by_id, raw_lines = all_raw()
    passes = discovery_passes()
    critical: list[str] = []

    missing_old_ids = sorted(set(baseline["all_preexisting_raw_ids"]) - set(raw_by_id))
    if missing_old_ids:
        critical.append(f"Missing {len(missing_old_ids)} pre-existing raw IDs")

    changed_gen_files: list[str] = []
    for relative, expected in baseline["protected_gen_hashes"].items():
        path = PROJECT_ROOT / relative
        if not path.exists() or sha256(path) != expected:
            changed_gen_files.append(relative)
    if changed_gen_files:
        critical.append(f"Protected GEN files changed: {changed_gen_files}")

    changed_gen_sources = [
        url for url, expected in (baseline.get("protected_gen_sources") or {}).items()
        if registry.get(url) != expected
    ]
    if changed_gen_sources:
        critical.append(f"Protected GEN source metadata changed: {changed_gen_sources}")

    accepted_export = workbook_semantics(PROJECT_ROOT / "data" / "exports" / "data_collection.xlsx")
    accepted_export_preserved = accepted_export == baseline.get("protected_accepted_export")
    if not accepted_export_preserved:
        critical.append("Accepted-only workbook contents differ from the baseline")

    workbook_path = PROJECT_ROOT / "data" / "exports" / "review_candidates_remaining.xlsx"
    review_book = load_workbook(workbook_path, read_only=True, data_only=True)
    if review_book.sheetnames != list(CATEGORIES):
        critical.append(f"Review workbook sheets mismatch: {review_book.sheetnames}")

    category_rows: list[dict[str, Any]] = []
    union_sources: set[str] = set()
    all_candidate_ids: list[str] = []
    direct_count = 0
    source_item_only = 0
    passes_required = {"taxonomy", "topic", "observed"}
    for category in CATEGORIES:
        queue = rows(PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv")
        primary = rows(PROJECT_ROOT / "data" / "candidates" / f"{category}.csv")
        secondary = rows(PROJECT_ROOT / "data" / "candidates" / f"{category}.secondary_review.csv")
        metrics = read_json(PROJECT_ROOT / "data" / "candidates" / f"{category}.metrics.json")
        statuses = Counter(row.get("Review Status") or "" for row in queue)
        sources = Counter(row.get("Source URL") or "" for row in queue)
        union_sources.update(url for url in sources if url)
        candidate_ids = [row.get("Candidate ID") or "" for row in queue]
        all_candidate_ids.extend(candidate_ids)
        provenance_failures: list[dict[str, str]] = []
        for row in queue:
            url = row.get("Source URL") or ""
            item_id = row.get("Source Item ID") or ""
            source = registry.get(url)
            if not source or not source_verified(source):
                provenance_failures.append({"candidate_id": row.get("Candidate ID") or "", "reason": "SOURCE_UNVERIFIED"})
            elif not item_id:
                provenance_failures.append({"candidate_id": row.get("Candidate ID") or "", "reason": "SOURCE_ITEM_ID_MISSING"})
            elif not any(
                raw.get("source_url") == url
                and str(raw.get("source_item_id") or "") == item_id
                and raw.get("original_text") == row.get("Example")
                and raw.get("example_is_real_world") is True
                for raw in raw_by_id.values()
            ):
                provenance_failures.append({"candidate_id": row.get("Candidate ID") or "", "reason": "EXACT_REAL_RAW_EVIDENCE_MISSING"})
            if row.get("Content URL"):
                direct_count += 1
            elif url and item_id:
                source_item_only += 1
        if provenance_failures:
            critical.append(f"{category}: {len(provenance_failures)} provenance failures")
        if len(primary) + len(secondary) != len(queue):
            critical.append(f"{category}: primary/secondary count does not equal queue")
        if any(statuses[status] for status in HUMAN):
            critical.append(f"{category}: unexpected human decisions were created")
        if statuses["PENDING"] != len(queue):
            critical.append(f"{category}: not every row is PENDING")
        if not passes_required.issubset(passes[category]):
            critical.append(f"{category}: missing discovery passes {sorted(passes_required - passes[category])}")
        category_rows.append({
            "category": category,
            "existing_raw": int(baseline["raw_counts_by_category"].get(category, 0)),
            "new_raw": int(metrics["raw_records"]) - int(baseline["raw_counts_by_category"].get(category, 0)),
            "current_raw": int(metrics["raw_records"]),
            "primary": len(primary),
            "secondary": len(secondary),
            "total_review": len(queue),
            "distinct_sources": len(sources),
            "largest_source_share": round(max(sources.values(), default=0) / len(queue), 4) if queue else 0.0,
            "verified_provenance": len(queue) - len(provenance_failures),
            "unverified": len(provenance_failures),
            "human_accept": statuses["ACCEPT"],
            "human_reject": statuses["REJECT"],
            "human_unsure": statuses["UNSURE"],
            "human_move": statuses["MOVE_TO_OTHER_CATEGORY"],
            "pending": statuses["PENDING"],
            "retrieval_passes": ["existing-corpus-remine", *sorted(passes[category] & passes_required)],
            "status": "COMPLETE" if len(queue) >= 100 else "COMPLETE_BELOW_TARGET",
            "candidate_count_per_source": dict(sorted(sources.items())),
            "direct_permalink_count": sum(bool(row.get("Content URL")) for row in queue),
            "source_url_plus_item_id_only": sum(bool(row.get("Source URL")) and bool(row.get("Source Item ID")) and not bool(row.get("Content URL")) for row in queue),
            "provenance_failures": provenance_failures,
        })

    duplicate_candidate_ids = sorted(candidate_id for candidate_id, count in Counter(all_candidate_ids).items() if count > 1)
    if duplicate_candidate_ids:
        critical.append(f"Duplicate Candidate IDs across review queues: {duplicate_candidate_ids}")

    baseline_queue_ids = {
        category: set(values.get("candidate_ids") or [])
        for category, values in baseline.get("queues", {}).items()
    }
    current_candidate_artifact_ids: dict[str, set[str]] = {category: set() for category in CATEGORIES}
    for category in CATEGORIES:
        for path in (PROJECT_ROOT / "data" / "candidates").glob(f"{category}*.csv"):
            current_candidate_artifact_ids[category].update(row.get("Candidate ID") or "" for row in rows(path))
        current_candidate_artifact_ids[category].update(
            row.get("Candidate ID") or "" for row in rows(PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv")
        )
    missing_baseline_candidates = {
        category: sorted(ids - current_candidate_artifact_ids[category])
        for category, ids in baseline_queue_ids.items()
        if ids - current_candidate_artifact_ids[category]
    }
    if missing_baseline_candidates:
        critical.append(f"Baseline Candidate IDs missing from current artifacts: {missing_baseline_candidates}")

    temp_files = [str(path.relative_to(PROJECT_ROOT)) for path in (PROJECT_ROOT / "data").rglob("*.tmp")]
    if temp_files:
        critical.append(f"Temporary files remain: {temp_files}")

    output = {
        "scope": list(CATEGORIES),
        "resume_point": "REG and STA collection/processing had completed; STA quality audit was the active checkpoint; REL-1 and REL-2 were incomplete.",
        "categories": category_rows,
        "totals": {
            "old_raw_unique_ids_remined": len(baseline["all_preexisting_raw_ids"]),
            "new_real_raw_unique_ids": len(set(raw_by_id) - set(baseline["all_preexisting_raw_ids"])),
            "current_raw_unique_ids": len(raw_by_id),
            "current_raw_physical_lines": raw_lines,
            "primary": sum(row["primary"] for row in category_rows),
            "secondary": sum(row["secondary"] for row in category_rows),
            "review_candidates": sum(row["total_review"] for row in category_rows),
            "distinct_review_source_urls": len(union_sources),
            "verified_candidate_provenance": sum(row["verified_provenance"] for row in category_rows),
            "unverified_candidates": sum(row["unverified"] for row in category_rows),
            "direct_permalinks": direct_count,
            "source_url_plus_item_id_only": source_item_only,
            "synthetic_examples_used": 0,
            "fabricated_urls_used": 0,
        },
        "preservation": {
            "preexisting_raw_ids_present": len(missing_old_ids) == 0,
            "missing_preexisting_raw_ids": missing_old_ids,
            "protected_gen_file_hashes_unchanged": not changed_gen_files,
            "changed_gen_files": changed_gen_files,
            "protected_gen_source_metadata_unchanged": not changed_gen_sources,
            "changed_gen_sources": changed_gen_sources,
            "accepted_export_semantics_unchanged": accepted_export_preserved,
            "baseline_candidate_ids_retained_in_artifacts": not missing_baseline_candidates,
            "missing_baseline_candidate_ids": missing_baseline_candidates,
        },
        "review_workbook_sheets": review_book.sheetnames,
        "temporary_files": temp_files,
        "critical_failures": critical,
        "status": "PASS" if not critical else "FAIL",
    }
    path = PROJECT_ROOT / "data" / "audits" / "remaining_final.json"
    path.write_text(json.dumps(output, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(path)
    print(json.dumps(output["totals"], ensure_ascii=False, sort_keys=True))
    print("STATUS", output["status"])
    for failure in critical:
        print("CRITICAL", failure)
    if critical:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


REMAINING = (
    "REG-1", "REG-2", "REG-3", "REG-4",
    "STA-1", "STA-2", "STA-3", "STA-4", "STA-5",
    "REL-1", "REL-2",
)
HUMAN_STATUSES = {"ACCEPT", "REJECT", "UNSURE", "MOVE_TO_OTHER_CATEGORY"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def protected_gen_files() -> list[Path]:
    roots = (
        PROJECT_ROOT / "data" / "raw",
        PROJECT_ROOT / "data" / "reviewed",
        PROJECT_ROOT / "data" / "rejected",
        PROJECT_ROOT / "data" / "candidates",
    )
    result: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        result.extend(path for path in root.rglob("GEN-*") if path.is_file())
    for relative in (
        "data/reviewed/final_ids.json",
        "config/taxonomy.json",
    ):
        path = PROJECT_ROOT / relative
        if path.exists():
            result.append(path)
    return sorted(set(result))


def workbook_semantics(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"exists": False, "sheets": {}}
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        header = [str(value or "") for value in values[0]] if values else []
        rows = [
            {header[index]: str(value or "") for index, value in enumerate(row)}
            for row in values[1:]
        ] if header else []
        sheets[sheet.title] = {"header": header, "rows": rows}
    return {"exists": True, "sheets": sheets}


def main() -> None:
    all_raw_ids: set[str] = set()
    raw_ids_by_category: dict[str, list[str]] = {}
    raw_counts_by_category: dict[str, int] = {}
    for path in sorted((PROJECT_ROOT / "data" / "raw").rglob("*.jsonl")):
        if path.name.endswith(".provenance.jsonl"):
            continue
        ids = {
            str(row.get("raw_record_id") or "")
            for row in read_jsonl(path)
            if row.get("raw_record_id")
        }
        all_raw_ids.update(ids)
        category = path.stem
        raw_ids_by_category.setdefault(category, []).extend(sorted(ids))
        raw_counts_by_category[category] = raw_counts_by_category.get(category, 0) + len(ids)

    queues: dict[str, Any] = {}
    for category in REMAINING:
        path = PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv"
        rows = csv_rows(path)
        queues[category] = {
            "sha256": sha256(path) if path.exists() else None,
            "candidate_ids": sorted(row.get("Candidate ID") or "" for row in rows),
            "human_decisions": {
                row.get("Candidate ID") or "": row.get("Review Status") or ""
                for row in rows
                if row.get("Review Status") in HUMAN_STATUSES
            },
            "row_count": len(rows),
        }


    registry = read_json(PROJECT_ROOT / "config" / "sources.json")
    sources_by_url = {str(row.get("url") or ""): row for row in registry.get("sources", [])}
    gen_source_urls: set[str] = set()
    for category in ("GEN-1", "GEN-2", "GEN-3", "GEN-4", "GEN-5"):
        for row in csv_rows(PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv"):
            if row.get("Source URL"):
                gen_source_urls.add(str(row["Source URL"]))
    protected_gen_sources = {
        url: sources_by_url.get(url)
        for url in sorted(gen_source_urls)
    }

    output = {
        "captured_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "scope": list(REMAINING),
        "protected_gen_hashes": {
            path.relative_to(PROJECT_ROOT).as_posix(): sha256(path)
            for path in protected_gen_files()
        },
        "protected_accepted_export": workbook_semantics(
            PROJECT_ROOT / "data" / "exports" / "data_collection.xlsx"
        ),
        "protected_gen_sources": protected_gen_sources,
        "source_registry_sha256": sha256(PROJECT_ROOT / "config" / "sources.json"),
        "all_preexisting_raw_ids": sorted(all_raw_ids),
        "raw_counts_by_category": raw_counts_by_category,
        "raw_ids_by_category": {
            category: sorted(set(ids)) for category, ids in raw_ids_by_category.items()
        },
        "queues": queues,
    }
    path = PROJECT_ROOT / "data" / "audits" / "remaining_baseline.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(output, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"BASELINE\t{path}")
    print(f"PROTECTED_GEN_FILES\t{len(output['protected_gen_hashes'])}")
    print(f"PREEXISTING_RAW_IDS\t{len(all_raw_ids)}")
    for category in REMAINING:
        queue = queues[category]
        print(
            f"{category}\tRAW={raw_counts_by_category.get(category, 0)}"
            f"\tQUEUE={queue['row_count']}\tHUMAN={len(queue['human_decisions'])}"
        )


if __name__ == "__main__":
    main()

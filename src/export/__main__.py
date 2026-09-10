from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


def workbook_path() -> Path:
    return PROJECT_ROOT / "data" / "exports" / "data_collection.xlsx"


def final_id_path() -> Path:
    return PROJECT_ROOT / "data" / "reviewed" / "final_ids.json"


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def load_evidence() -> tuple[dict[str, dict], dict[str, list[dict]]]:
    raw_by_id: dict[str, dict] = {}
    for path in (PROJECT_ROOT / "data" / "raw").rglob("*.jsonl"):
        # Provenance sidecars reuse raw_record_id but are not raw comment records;
        # loading them here would overwrite the actual text evidence.
        if path.name.endswith(".provenance.jsonl"):
            continue
        for record in read_jsonl(path):
            if record.get("raw_record_id") and record.get("original_text"):
                raw_by_id[record["raw_record_id"]] = record
    occurrences_by_candidate: dict[str, list[dict]] = {}
    for path in (PROJECT_ROOT / "data" / "candidates").glob("*.occurrences.jsonl"):
        for occurrence in read_jsonl(path):
            candidate_id = occurrence.get("candidate_id")
            if candidate_id:
                occurrences_by_candidate.setdefault(candidate_id, []).append(occurrence)
    return raw_by_id, occurrences_by_candidate


def load_current_review_rows(code: str) -> list[dict[str, str]]:
    queue_path = PROJECT_ROOT / "data" / "reviewed" / f"{code}.review_queue.csv"
    if queue_path.exists():
        return read_csv(queue_path)
    return read_csv(PROJECT_ROOT / "data" / "reviewed" / f"{code}.csv")


def current_exportable_rows(
    code: str,
    verified_urls: set[str],
    raw_by_id: dict[str, dict],
    occurrences: dict[str, list[dict]],
) -> list[dict[str, str]]:
    return [
        row
        for row in load_current_review_rows(code)
        if eligible(row, verified_urls, raw_by_id, occurrences)
    ]


def source_is_verified(source: dict) -> bool:
    """Require reachability and identity evidence, not a legacy HTTP-only flag."""
    if not source.get("public_access"):
        return False
    reachable = source.get("reachability_verified")
    if reachable is None:
        reachable = source.get("reachability", {}).get("reachable")
    identity = source.get("identity_verified")
    if identity is None:
        identity = source.get("identity_verification", {}).get("verified")
    return bool(reachable and identity)


def eligible(row: dict[str, str], verified_urls: set[str], raw_by_id: dict[str, dict], occurrences: dict[str, list[dict]]) -> bool:
    if row.get("Review Status") != "ACCEPT" or not row.get("Example"):
        return False
    if row.get("Source URL") not in verified_urls or not row.get("Source Item ID"):
        return False
    evidence = []
    for occurrence in occurrences.get(row.get("Candidate ID", ""), []):
        raw = raw_by_id.get(str(occurrence.get("raw_record_id")))
        if raw:
            evidence.append(raw)
    return any(
        record.get("example_is_real_world") is True
        and record.get("source_verified") is True
        and record.get("original_text") == row.get("Example")
        for record in evidence
    )


def allocate_id(code: str, candidate_id: str, mapping: dict[str, str]) -> str:
    key = f"{code}:{candidate_id}"
    if key in mapping:
        return mapping[key]
    prefix = code.replace("-", "")
    used = [int(value.rsplit("-", 1)[1]) for value in mapping.values() if value.startswith(prefix + "-")]
    mapping[key] = f"{prefix}-{(max(used, default=0) + 1):03d}"
    return mapping[key]


def main() -> None:
    taxonomy = read_json(PROJECT_ROOT / "config" / "taxonomy.json")
    registry = read_json(PROJECT_ROOT / "config" / "sources.json")
    verified_urls = {source["url"] for source in registry["sources"] if source_is_verified(source)}
    raw_by_id, occurrences = load_evidence()
    id_path = final_id_path()
    mapping = read_json(id_path) if id_path.exists() else {}

    workbook = Workbook()
    workbook.remove(workbook.active)
    exported = 0
    for subcategory in taxonomy["subcategories"]:
        code = subcategory["code"]
        sheet = workbook.create_sheet(title=code[:31])
        headers = ["Example ID", "Example", "Context", "Source"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:D1"

        reviewed_rows = current_exportable_rows(code, verified_urls, raw_by_id, occurrences)
        for row in reviewed_rows:
            if not eligible(row, verified_urls, raw_by_id, occurrences):
                continue
            example_id = allocate_id(code, row["Candidate ID"], mapping)
            link = row.get("Content URL") or row.get("Source URL")
            if row.get("Content URL"):
                source_label = f"{row.get('Source Platform', 'Public source')} — original item"
            else:
                platform = row.get("Source Platform") or "Public source"
                noun = "video" if platform == "YouTube" else "source page"
                source_label = f"{platform} {noun} — direct comment permalink unavailable"
            sheet.append([example_id, row["Example"], row.get("Context", ""), source_label])
            source_cell = sheet.cell(row=sheet.max_row, column=4)
            source_cell.hyperlink = link
            source_cell.style = "Hyperlink"
            exported += 1

        for row_cells in sheet.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        widths = {1: 14, 2: 70, 3: 55, 4: 45}
        for column, width in widths.items():
            sheet.column_dimensions[get_column_letter(column)].width = width

    output = workbook_path()
    output.parent.mkdir(parents=True, exist_ok=True)
    id_path.parent.mkdir(parents=True, exist_ok=True)
    id_path.write_text(json.dumps(mapping, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    workbook.save(output)
    print(f"EXPORTED\t{exported}")
    print(f"WORKBOOK\t{output}")


if __name__ == "__main__":
    main()

from __future__ import annotations

from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.export.consolidation import enriched_rows, source_registry
from src.utils.io import PROJECT_ROOT


CATEGORIES = ("GEN-1", "REG-1", "REG-3", "REG-4")
OUTPUT = PROJECT_ROOT / "data" / "exports" / "gap_fill_review_batch_A.xlsx"
COLUMNS = [
    "Example", "Context", "Source Link", "Candidate ID", "Source Item ID",
    "Human-readable Provenance", "Retrieval Tier", "Retrieval Class",
    "Review Priority Score", "Stance Hint", "Suggested Category", "Review Status",
]


def main() -> None:
    registry = source_registry()
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts: Counter[str] = Counter()
    for category in CATEGORIES:
        _, rows = enriched_rows(category, registry)
        pending = [row for row in rows if row.get("Review Status") == "PENDING"]
        pending.sort(key=lambda row: (-int(row.get("Review Priority Score") or 0), row.get("Candidate ID") or ""))
        counts[category] = len(pending)
        sheet = workbook.create_sheet(category)
        sheet.append(COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in pending:
            sheet.append([row.get(column, "") for column in COLUMNS])
            link = sheet.cell(sheet.max_row, COLUMNS.index("Source Link") + 1)
            if isinstance(link.value, str) and link.value.startswith(("https://", "http://")):
                link.hyperlink = link.value
                link.style = "Hyperlink"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, sheet.max_row)}"
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, column in enumerate(COLUMNS, start=1):
            width = 72 if column == "Example" else 55 if column in {"Context", "Human-readable Provenance"} else 28
            sheet.column_dimensions[get_column_letter(index)].width = width
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT.with_name(OUTPUT.stem + ".tmp" + OUTPUT.suffix)
    workbook.save(temporary)
    temporary.replace(OUTPUT)
    print(OUTPUT)
    print(dict(counts))


if __name__ == "__main__":
    main()

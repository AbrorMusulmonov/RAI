"""Build the human-review workbook for REG/STA/REL PENDING queues."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.utils.io import PROJECT_ROOT


CATEGORIES = (
    "REG-1", "REG-2", "REG-3", "REG-4",
    "STA-1", "STA-2", "STA-3", "STA-4", "STA-5",
    "REL-1", "REL-2",
)
HEADERS = (
    "Example", "Context", "Source Link", "Candidate ID", "Retrieval Tier",
    "Stance Hint", "Suggested Category", "Possible Compound Category",
)


def read_queue(category: str) -> list[dict[str, str]]:
    path = PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv"
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def main() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    total = 0
    for category in CATEGORIES:
        sheet = workbook.create_sheet(category)
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
        for row in read_queue(category):
            # This artifact mirrors the authoritative queue; it never changes a
            # review status and never treats a model suggestion as acceptance.
            source_link = row.get("Content URL") or row.get("Source URL") or ""
            sheet.append([
                row.get("Example") or "",
                row.get("Context") or "",
                source_link,
                row.get("Candidate ID") or "",
                row.get("Retrieval Tier") or "",
                row.get("Stance Hint") or "",
                row.get("Suggested Category") or "",
                row.get("Possible Compound Category") or "",
            ])
            link_cell = sheet.cell(row=sheet.max_row, column=3)
            if source_link:
                link_cell.hyperlink = source_link
                link_cell.style = "Hyperlink"
            total += 1
        widths = (72, 58, 48, 24, 18, 26, 20, 28)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    output = PROJECT_ROOT / "data" / "exports" / "review_candidates_remaining.xlsx"
    temporary = output.with_suffix(".xlsx.tmp")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(temporary)
    temporary.replace(output)
    print(f"EXPORTED_REVIEW_CANDIDATES\t{total}")
    print(f"WORKBOOK\t{output}")


if __name__ == "__main__":
    main()

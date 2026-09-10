"""Canonical, formula-driven Excel interface. Exporting never annotates data."""
from __future__ import annotations
import json
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter
from src.audit.final50 import ROOT, AUDIT, CATEGORIES, HUMAN, queues, profile, save, now
from src.processing.category import atomic_csv

STATUSES=('PENDING','ACCEPT','REJECT','UNSURE','MOVE_TO_OTHER_CATEGORY')
COLUMNS=['Candidate ID','Example','Context','Source Link','Source Platform','Source Name','Source Item ID',
 'Retrieval Tier','Review Priority Score','Stance Hint','Suggested Category','Possible Compound Category','Review Status','Reviewer Notes']
DASH_COLUMNS=['Category','Total Review Candidates','PRIMARY','SECONDARY','ACCEPT','REJECT','UNSURE','MOVE','PENDING',
 'Distinct Sources','Estimated Review Precision','Estimated Accepted Yield','Gap to 50','Collection Readiness']
EXPORTS=ROOT/'data/exports'

def excel_text(cell,value):
    """Force public text to literal strings, including leading =,+,-,@."""
    cell.value=str(value or '');cell.data_type='s'

def score(r):
    try: return min(100,max(0,float(r.get('Review Priority Score') or (80 if r.get('Retrieval Tier')=='PRIMARY' else 55))))
    except (TypeError,ValueError): return 55

def ordered(rows):
    reviewed=[r for r in rows if r['Review Status']!='PENDING']
    pending=sorted((r for r in rows if r['Review Status']=='PENDING'),key=lambda r:(-score(r),r.get('Retrieval Tier')!='PRIMARY',r['Candidate ID']))
    # Round-robin within the same priority/tier preserves likely relevance first.
    result=[]
    while pending:
        priority=(score(pending[0]),pending[0].get('Retrieval Tier'));seen=set()
        for r in list(pending):
            if (score(r),r.get('Retrieval Tier'))==priority and r['Source URL'] not in seen:
                result.append(r);pending.remove(r);seen.add(r['Source URL'])
    return reviewed+result

def style_header(ws,widths):
    ws.freeze_panes='B2';ws.sheet_view.showGridLines=False;ws.row_dimensions[1].height=34
    for cell in ws[1]:
        cell.font=Font(name='Calibri',bold=True,color='FFFFFF',size=11)
        cell.fill=PatternFill('solid',fgColor='183D52');cell.alignment=Alignment(wrap_text=True,vertical='center')
    for i,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=width

def review_sheet(w,c,rows):
    ws=w.create_sheet(c);ws.append(COLUMNS)
    for r in ordered(rows):
        assert r['Review Status'] in STATUSES
        values={**r,'Source Link':r.get('Content URL') or r['Source URL'],'Review Priority Score':score(r),
                'Suggested Category':r.get('Suggested Category') or c}
        index=ws.max_row+1
        for j,k in enumerate(COLUMNS,1):
            cell=ws.cell(index,j)
            if k=='Review Priority Score':cell.value=values[k]
            else:excel_text(cell,values.get(k,''))
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            cell.font=Font(name='Calibri',size=11)
        link=ws.cell(index,4);link.hyperlink=values['Source Link'];link.style='Hyperlink'
        ws.row_dimensions[index].height=min(240,max(50,15*(1+len(r['Example'])//70)))
    style_header(ws,[26,80,49,43,17,34,34,16,18,26,20,23,28,45])
    last=max(1,ws.max_row);ws.auto_filter.ref=f'A1:N{last}'
    if rows:
        t=Table(displayName='Review_'+c.replace('-','_'),ref=ws.auto_filter.ref)
        t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False)
        ws.add_table(t)
    dv=DataValidation(type='list',formula1='"'+','.join(STATUSES)+'"',allow_blank=False)
    dv.errorTitle='Choose a review status';dv.error='Use one of the five dropdown choices.'
    dv.showErrorMessage=True;dv.errorStyle='stop';dv.showInputMessage=True
    dv.promptTitle='Human review';dv.prompt='Only you decide. Keep Candidate ID and Example unchanged.'
    ws.add_data_validation(dv);dv.add(f'M2:M{max(1000,last)}')
    for status,color in [('ACCEPT','DDF2DF'),('REJECT','FCE0E0'),('UNSURE','FFF2C7'),('MOVE_TO_OTHER_CATEGORY','DDEDFC')]:
        ws.conditional_formatting.add(f'M2:M{max(2,last)}',FormulaRule(formula=[f'$M2="{status}"'],fill=PatternFill('solid',fgColor=color)))
    return ws

def assert_no_unmerged_edits(path,current):
    if not path.exists():return
    w=load_workbook(path,read_only=True)
    try:
        for c in CATEGORIES:
            if c not in w.sheetnames:continue
            it=w[c].iter_rows(values_only=True);headers=next(it)
            if 'Candidate ID' not in headers or 'Review Status' not in headers:continue
            existing={r['Candidate ID']:r for r in current[c]}
            for values in it:
                row=dict(zip(headers,values));cid=row.get('Candidate ID')
                if not cid:continue
                canonical=existing.get(cid)
                if row.get('Review Status') in HUMAN and (not canonical or canonical['Review Status']!=row['Review Status']):
                    raise RuntimeError(f'Unmerged human decision in {path.name}: {cid}. Refusing overwrite.')
                if (row.get('Reviewer Notes') or '') and (not canonical or (canonical.get('Reviewer Notes') or '')!=row['Reviewer Notes']):
                    raise RuntimeError(f'Unmerged reviewer notes: {cid}. Refusing overwrite.')
    finally:w.close()

def write_main(current,matrix,path):
    assert_no_unmerged_edits(path,current)
    w=Workbook();dash=w.active;dash.title='DASHBOARD';dash.append(DASH_COLUMNS)
    ins=w.create_sheet('INSTRUCTIONS');ins.append(['Topic','Instructions'])
    messages=[
      ('Goal','Har kategoriya uchun taxminan 50 ta inson ACCEPT qarori. Hozirgi fayl — nomzodlar, tayyor dataset emas.'),
      ('1. Read','Kategoriya varag‘ini oching. Example va Context ni o‘qing. RAI_in_Full.pdf ta’riflariga solishtiring.'),
      ('2. Check source','Source Link ni bosing. To‘g‘ridan-to‘g‘ri comment havolasi bo‘lmasa video ochiladi; aniq izoh IDsi Source Item ID ustunida.'),
      ('3. Choose','Review Status ustunidagi dropdown orqali qaror tanlang. Reviewer Notes ga izoh yoki ko‘chiriladigan kategoriyani yozing.'),
      ('PENDING','Hali inson ko‘rib chiqmagan. Model score/tier yakuniy baho emas.'),
      ('ACCEPT','Misol shu kategoriyaning to‘liq ta’rifiga mos keladi.'),
      ('REJECT','Mos emas: neytral, noto‘g‘ri kategoriya, shovqin yoki boshqa sabab.'),
      ('UNSURE','Keyinroq qayta ko‘rib chiqish kerak.'),
      ('MOVE_TO_OTHER_CATEGORY','Foydali misol, lekin boshqa kategoriya mosroq. Uning kodini Reviewer Notes ga yozing.'),
      ('Do not edit','Candidate ID va original Example ni O‘ZGARTIRMANG. Barcha qatorlarni saqlang; filtrdan foydalaning.'),
      ('Save','Excel faylini saqlang. Qarorlar keyinchalik Candidate ID bilan ichki CSV navbatlariga birlashtiriladi. Excel’ni saqlash CSV ni avtomatik o‘zgartirmaydi.'),
      ('No terminal','Oddiy inson review uchun terminal yoki CLI buyrug‘i kerak emas. Shu Excel fayli bilan ishlang.'),
      ('Dashboard','Status sonlari Excel’da avtomatik hisoblanadi. Estimated Accepted Yield = ACCEPT + PENDING × estimated precision; bu kafolat emas. Gap to 50 — hali inson ACCEPT qilishi kerak bo‘lgan son.'),
      ('Snapshot estimates','Distinct Sources, Estimated Review Precision va Collection Readiness — eksport paytidagi audit. Ular inson qarorlari emas; formulalar Excel ochilganda qayta hisoblanadi.'),
      ('Safety','Original izohlar haqoratli bo‘lishi mumkin. Matnlar real raw yozuvlardan; AI faqat topish va saralashga yordam bergan.'),
      ('Final dataset','data_collection.xlsx va aziza_accepted.xlsx faqat avval inson ACCEPT qilgan misollarni o‘z ichiga oladi.'),
    ]
    for item in messages:ins.append(item)
    style_header(ins,[25,115]);ins.freeze_panes='A2'
    for row in ins.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        ins.row_dimensions[row[0].row].height=48
    for i,c in enumerate(CATEGORIES,2):
        review_sheet(w,c,current[c]);end=max(2,len(current[c])+1);ref=f"'{c}'!"
        statuses=ref+f'$M$2:$M${end}';tier=ref+f'$H$2:$H${end}'
        dash.append([c,f'=COUNTA({ref}$A$2:$A${end})',f'=COUNTIF({tier},"PRIMARY")',f'=COUNTIF({tier},"SECONDARY")',
          *[f'=COUNTIF({statuses},"{s}")' for s in ('ACCEPT','REJECT','UNSURE','MOVE_TO_OTHER_CATEGORY','PENDING')],
          matrix[c]['Distinct Sources'],matrix[c]['Estimated Review Precision'],f'=E{i}+I{i}*K{i}',f'=MAX(0,50-E{i})',matrix[c]['Readiness Status']])
        dash.cell(i,1).hyperlink=f"#'{c}'!A1";dash.cell(i,1).style='Hyperlink';dash.cell(i,11).number_format='0%';dash.cell(i,12).number_format='0.0'
    style_header(dash,[15,22,14,14,12,12,12,13,13,18,22,23,16,31]);dash.auto_filter.ref='A1:N17'
    for row in dash.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        dash.row_dimensions[row[0].row].height=32
    w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    path.parent.mkdir(parents=True,exist_ok=True);temp=path.with_suffix('.tmp.xlsx');w.save(temp);temp.replace(path)

def export_all():
    current=queues();matrix=profile()
    if any(r['Provenance Failures'] for r in matrix.values()):raise RuntimeError('Active queue provenance failure')
    write_main(current,matrix,EXPORTS/'FINAL_HUMAN_REVIEW.xlsx')
    top=Workbook();top.remove(top.active)
    for c in CATEGORIES:review_sheet(top,c,[r for r in ordered(current[c]) if r['Review Status']=='PENDING'][:30])
    assert_no_unmerged_edits(EXPORTS/'top_candidates_for_review.xlsx',current)
    top.save(EXPORTS/'top_candidates_for_review.xlsx')
    accepted=Workbook();accepted.remove(accepted.active)
    for c in CATEGORIES:
        ws=accepted.create_sheet(c);ws.append(['Example ID','Example','Context','Source Link'])
        for r in current[c]:
            if r['Review Status']!='ACCEPT':continue
            n=ws.max_row+1
            for j,v in enumerate(['',r['Example'],r.get('Context',''),r.get('Content URL') or r['Source URL']],1):excel_text(ws.cell(n,j),v)
            ws.cell(n,4).hyperlink=ws.cell(n,4).value;ws.cell(n,4).style='Hyperlink'
        style_header(ws,[18,80,55,50]);ws.auto_filter.ref=f'A1:D{ws.max_row}'
    accepted.save(EXPORTS/'aziza_accepted.xlsx')
    fields=['Category','Raw Records','Unique Raw','Primary','Secondary','Total Review','Human ACCEPT','Human REJECT',
      'Human UNSURE','Human MOVE','PENDING','Distinct Sources','Largest Source Share','Verified Provenance',
      'Estimated Review Precision','Estimated Accepted Yield','Gap to 50','Readiness Status','Main Limitation']
    atomic_csv(EXPORTS/'category_readiness_summary.csv',fields,[{k:r[k] for k in fields} for r in matrix.values()])
    save('current.json',matrix)
    status=Counter(r['Review Status'] for rows in current.values() for r in rows)
    save('excel_manifest.json',{'timestamp':now(),'category_sheets':16,'total_rows':sum(map(len,current.values())),
        'statuses':dict(status),'distinct_source_urls':len({r['Source URL'] for rows in current.values() for r in rows}),
        'queue_ids':{c:[r['Candidate ID'] for r in ordered(rows)] for c,rows in current.items()}})
    print(json.dumps({'path':str(EXPORTS/'FINAL_HUMAN_REVIEW.xlsx'),'statuses':dict(status)}))

if __name__=='__main__':export_all()

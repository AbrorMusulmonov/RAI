"""Create a read-only snapshot workbook of every current review queue."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.utils.io import PROJECT_ROOT, read_json


OUTPUT = PROJECT_ROOT / "data" / "exports" / "draft_all_subcategories.xlsx"
COLUMNS = [
    "Category", "Candidate ID", "Example", "Context", "Source Platform", "Source Name",
    "Source Link", "Source URL", "Content URL", "Source Item ID", "Human-readable Provenance",
    "Retrieval Tier", "Retrieval Class", "Retrieval Signal", "Stance Hint", "Stance Evidence",
    "Model Retrieval Relevance", "Model Evidence Span", "Suggested Category",
    "Possible Compound Category", "Review Priority Score", "Review Status", "Reviewer Notes",
]


def read_queue(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def main() -> None:
    taxonomy = read_json(PROJECT_ROOT / "config" / "taxonomy.json")
    categories = [item["code"] for item in taxonomy["subcategories"]]
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts = {}
    for category in categories:
        queue = read_queue(PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv")
        counts[category] = len(queue)
        sheet = workbook.create_sheet(category)
        sheet.append(COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for source_row in queue:
            row = dict(source_row)
            source_url = row.get("Source URL") or ""
            item_id = row.get("Source Item ID") or ""
            row["Category"] = category
            row["Source Link"] = row.get("Content URL") or source_url
            if not row.get("Human-readable Provenance"):
                row["Human-readable Provenance"] = (
                    f"{row.get('Source Name') or ''} | {source_url} | Source Item ID: {item_id}"
                )
            sheet.append([row.get(column, "") for column in COLUMNS])
            link_cell = sheet.cell(sheet.max_row, COLUMNS.index("Source Link") + 1)
            if isinstance(link_cell.value, str) and link_cell.value.startswith(("https://", "http://")):
                link_cell.hyperlink = link_cell.value
                link_cell.style = "Hyperlink"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, sheet.max_row)}"
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, column in enumerate(COLUMNS, start=1):
            width = 72 if column == "Example" else 52 if column in {"Context", "Retrieval Signal", "Human-readable Provenance"} else 25
            sheet.column_dimensions[get_column_letter(index)].width = width
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT.with_name(OUTPUT.stem + ".tmp" + OUTPUT.suffix)
    workbook.save(temporary)
    temporary.replace(OUTPUT)
    print(OUTPUT)
    print(counts)


if __name__ == "__main__":
    main()

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict, deque
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.review.__main__ import atomic_write, load_rows
from src.utils.io import PROJECT_ROOT, read_json, read_jsonl


EXPORT_DIR = PROJECT_ROOT / "data" / "exports"
MASTER_PATH = EXPORT_DIR / "master_human_review.xlsx"
TOP_PATH = EXPORT_DIR / "top_candidates_for_review.xlsx"
READINESS_PATH = EXPORT_DIR / "category_readiness_summary.csv"
SOURCE_COUNTS_PATH = EXPORT_DIR / "candidate_counts_per_source.csv"
BASELINE_PATH = PROJECT_ROOT / "data" / "audits" / "consolidation_baseline.json"
PROFILE_PATH = PROJECT_ROOT / "data" / "audits" / "consolidation_profile.json"
MANUAL_PATH = PROJECT_ROOT / "data" / "audits" / "consolidation_manual_assessments.json"
READINESS_CONFIG_PATH = PROJECT_ROOT / "config" / "consolidation_readiness.json"

REVIEW_COLUMNS = [
    "Candidate ID", "Example", "Context", "Source Platform", "Source Name",
    "Source Link", "Source URL", "Content URL", "Source Item ID", "Human-readable Provenance",
    "Provenance Status", "Retrieval Tier", "Retrieval Class", "Retrieval Signal",
    "Stance Hint", "Stance Evidence", "Model Retrieval Relevance",
    "Model Suggested Category", "Model Evidence Span", "Model Stance Hint",
    "Model Confidence", "Suggested Category", "Possible Compound Category",
    "Review Priority Score", "Priority Basis", "Review Status", "Reviewer Notes",
]
TOP_COLUMNS = [
    "Balanced Review Rank", "Example", "Context", "Source Link", "Candidate ID",
    "Source Item ID", "Retrieval Tier", "Review Priority Score", "Stance Hint",
    "Possible Compound Category", "Review Status", "Reviewer Notes",
]


def float_value(value: Any, default: float = 0.0) -> float:
    try:
        return float(value or default)
    except (TypeError, ValueError):
        return default


def source_registry() -> dict[str, dict[str, Any]]:
    registry = read_json(PROJECT_ROOT / "config" / "sources.json")
    return {str(source.get("url") or ""): source for source in registry["sources"]}


def source_verified(source: dict[str, Any] | None) -> bool:
    if not source or not source.get("public_access"):
        return False
    reachable = source.get("reachability", {}).get("reachable")
    if reachable is None:
        reachable = source.get("reachability_verified")
    identity = source.get("identity_verification", {}).get("verified")
    if identity is None:
        identity = source.get("identity_verified")
    return bool(reachable and identity)


def upstream_retrieval_tiers(category: str) -> dict[str, str]:
    tiers: dict[str, str] = {}
    for suffix in (".csv", ".secondary_review.csv", ".broad_discovery.csv"):
        path = PROJECT_ROOT / "data" / "candidates" / f"{category}{suffix}"
        _, rows = load_rows(path)
        for row in rows:
            candidate_id = row.get("Candidate ID") or ""
            if candidate_id and candidate_id not in tiers:
                tiers[candidate_id] = row.get("Retrieval Tier") or ""
    return tiers


def retrieval_class(row: dict[str, str], upstream_tier: str) -> str:
    signal = (row.get("Retrieval Signal") or "").casefold()
    if "strong lexical signal" in signal or upstream_tier == "STRONG_LEXICAL":
        return "STRONG_LEXICAL"
    return "SEMANTIC_CONTEXTUAL"


def priority(row: dict[str, str], retrieval: str, verified: bool) -> tuple[int, str]:
    score = 25.0
    reasons: list[str] = ["base=25"]
    tier = (row.get("Retrieval Tier") or "").upper()
    if tier in {"PRIMARY", "STRONG_LEXICAL"}:
        score += 25
        reasons.append("primary_or_strong=+25")
    else:
        score += 12
        reasons.append("secondary_or_contextual=+12")
    if retrieval == "STRONG_LEXICAL":
        score += 15
        reasons.append("strong_lexical=+15")
    else:
        relevance = float_value(row.get("Model Retrieval Relevance"))
        if relevance:
            addition = min(15.0, relevance * 15.0)
            score += addition
            reasons.append(f"semantic_relevance=+{addition:.1f}")
        else:
            score += 7
            reasons.append("contextual_support=+7")
    stance = row.get("Stance Hint") or ""
    stance_adjustment = {
        "POSSIBLE_ABUSE": 12,
        "AMBIGUOUS": 6,
        "QUOTED_OR_REPORTED_ABUSE": 2,
        "POSSIBLE_COUNTER_SPEECH": 0,
    }.get(stance, 3)
    score += stance_adjustment
    reasons.append(f"stance_{stance or 'UNKNOWN'}=+{stance_adjustment}")
    if verified and row.get("Source Item ID"):
        score += 5
        reasons.append("verified_provenance=+5")
    if row.get("Possible Compound Category"):
        score += 3
        reasons.append("compound_review_value=+3")
    return int(round(max(0, min(100, score)))), "; ".join(reasons)


def enriched_rows(category: str, registry: dict[str, dict[str, Any]]) -> tuple[list[str], list[dict[str, str]]]:
    queue_path = PROJECT_ROOT / "data" / "reviewed" / f"{category}.review_queue.csv"
    fields, rows = load_rows(queue_path)
    upstream = upstream_retrieval_tiers(category)
    enriched: list[dict[str, str]] = []
    for row in rows:
        result = dict(row)
        source_url = result.get("Source URL") or ""
        source_item_id = result.get("Source Item ID") or ""
        verified = source_verified(registry.get(source_url))
        klass = retrieval_class(result, upstream.get(result.get("Candidate ID") or "", ""))
        score, basis = priority(result, klass, verified)
        result["Human-readable Provenance"] = (
            f"{result.get('Source Name') or ''} | {source_url} | Source Item ID: {source_item_id}"
        )
        result["Source Link"] = result.get("Content URL") or source_url
        result["Provenance Status"] = (
            "VERIFIED_SOURCE_DIRECT_PERMALINK" if verified and result.get("Content URL")
            else "VERIFIED_SOURCE_URL_PLUS_ITEM_ID" if verified and source_item_id
            else "PROVENANCE_FAILURE"
        )
        result["Retrieval Class"] = klass
        result["Review Priority Score"] = str(score)
        result["Priority Basis"] = basis
        enriched.append(result)

    reviewed = [row for row in enriched if row.get("Review Status") != "PENDING"]
    pending = [row for row in enriched if row.get("Review Status") == "PENDING"]
    pending.sort(key=lambda row: (-int(row["Review Priority Score"]), row.get("Candidate ID") or ""))
    ordered = reviewed + pending
    queue_fields = list(dict.fromkeys(fields + [
        "Human-readable Provenance", "Provenance Status", "Retrieval Class",
        "Review Priority Score", "Priority Basis",
    ]))
    atomic_write(queue_path, queue_fields, ordered)
    return queue_fields, ordered


def balanced_pending(rows: list[dict[str, str]], limit: int) -> list[dict[str, str]]:
    by_source: dict[str, deque[dict[str, str]]] = defaultdict(deque)
    for row in sorted(rows, key=lambda item: (-int(item["Review Priority Score"]), item.get("Candidate ID") or "")):
        by_source[row.get("Source URL") or "MISSING"].append(row)
    selected: list[dict[str, str]] = []
    source_uses: Counter[str] = Counter()
    while by_source and len(selected) < limit:
        source = max(
            by_source,
            key=lambda value: (
                int(by_source[value][0]["Review Priority Score"]) - 8 * source_uses[value],
                -source_uses[value],
                value,
            ),
        )
        row = by_source[source].popleft()
        row = dict(row)
        row["Balanced Review Rank"] = str(len(selected) + 1)
        selected.append(row)
        source_uses[source] += 1
        if not by_source[source]:
            del by_source[source]
    return selected


def write_workbook(path: Path, category_rows: dict[str, list[dict[str, str]]], top: bool = False) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    columns = TOP_COLUMNS if top else REVIEW_COLUMNS
    for category, rows in category_rows.items():
        sheet = workbook.create_sheet(category)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in rows:
            sheet.append([row.get(column, "") for column in columns])
            for link_column in ("Source Link", "Source URL", "Content URL"):
                if link_column not in columns:
                    continue
                column_index = columns.index(link_column) + 1
                cell = sheet.cell(row=sheet.max_row, column=column_index)
                if isinstance(cell.value, str) and cell.value.startswith(("https://", "http://")):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, sheet.max_row)}"
        for row_cells in sheet.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, column in enumerate(columns, start=1):
            width = 70 if column == "Example" else 55 if column in {"Context", "Retrieval Signal", "Stance Evidence", "Priority Basis"} else 28
            sheet.column_dimensions[get_column_letter(index)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".tmp" + path.suffix)
    workbook.save(temporary)
    temporary.replace(path)


def raw_counts() -> tuple[dict[str, int], dict[str, int]]:
    physical: dict[str, int] = {}
    unique: dict[str, int] = {}
    for path in sorted((PROJECT_ROOT / "data" / "raw" / "youtube").glob("*.jsonl")):
        if path.name.endswith(".provenance.jsonl"):
            continue
        records = read_jsonl(path)
        physical[path.stem] = len(records)
        unique[path.stem] = len({record.get("raw_record_id") for record in records if record.get("raw_record_id")})
    return physical, unique


def write_readiness(
    category_rows: dict[str, list[dict[str, str]]], profile: dict[str, Any], baseline: dict[str, Any],
    manual: dict[str, Any], readiness: dict[str, Any],
) -> None:
    fields = [
        "Category", "Raw Records", "Unique Raw", "Primary", "Secondary", "Total Review",
        "Human ACCEPT", "Human REJECT", "Human UNSURE", "Human MOVE", "PENDING",
        "Distinct Sources", "Largest Source Share", "Verified Provenance",
        "Estimated Review Precision", "Estimated Accepted Yield", "Readiness Status", "Main Limitation",
        "Unique Raw Records", "Gap-fill New Raw Records",
        "Candidates Before Precision Filtering", "Quality-excluded PENDING", "Candidates After Precision Filtering",
        "Strong Lexical", "Semantic/Contextual", "Direct Permalink", "Source URL + Platform Item ID Only",
        "ACCEPT", "REJECT", "UNSURE", "MOVE_TO_OTHER_CATEGORY",
        "Estimated Precision Before Cleanup", "Estimated Precision After Cleanup",
        "Health", "Readiness Group", "Main Issue", "Recommendation",
    ]
    physical, unique = raw_counts()
    rows: list[dict[str, Any]] = []
    for category, current in category_rows.items():
        statuses = Counter(row.get("Review Status") or "" for row in current)
        tiers = Counter(row.get("Retrieval Tier") or "" for row in current)
        classes = Counter(row.get("Retrieval Class") or "" for row in current)
        before = len(baseline["queues"][category]["candidate_ids"])
        profile_row = profile["categories"][category]
        manual_row = manual["categories"][category]
        ready_row = readiness["categories"][category]
        readiness_status = (
            "READY_FOR_HUMAN_REVIEW"
            if ready_row["readiness_group"] == "READY_FOR_HUMAN_REVIEW"
            else "SHORTFALL_DOCUMENTED"
            if ready_row["readiness_group"] == "BLOCKED_OR_SYSTEMATIC_RETRIEVAL_PROBLEM"
            else "READY_BUT_SMALL"
        )
        rows.append({
            "Category": category,
            "Raw Records": physical.get(category, 0),
            "Unique Raw": unique.get(category, 0),
            "Total Review": len(current),
            "Human ACCEPT": statuses["ACCEPT"],
            "Human REJECT": statuses["REJECT"],
            "Human UNSURE": statuses["UNSURE"],
            "Human MOVE": statuses["MOVE_TO_OTHER_CATEGORY"],
            "Verified Provenance": profile_row["verified_provenance"],
            "Estimated Review Precision": ready_row["estimated_precision_after_cleanup"],
            "Readiness Status": readiness_status,
            "Main Limitation": manual_row["main_issue"],
            "Unique Raw Records": unique.get(category, 0),
            "Gap-fill New Raw Records": physical.get(category, 0) - baseline["raw_counts_by_file_stem"].get(category, 0),
            "Candidates Before Precision Filtering": before,
            "Quality-excluded PENDING": before - len(current),
            "Candidates After Precision Filtering": len(current),
            "Primary": tiers["PRIMARY"],
            "Secondary": tiers["SECONDARY"],
            "Strong Lexical": classes["STRONG_LEXICAL"],
            "Semantic/Contextual": classes["SEMANTIC_CONTEXTUAL"],
            "Distinct Sources": profile_row["distinct_source_urls"],
            "Largest Source Share": profile_row["largest_source_share"],
            "Direct Permalink": profile_row["direct_permalink_count"],
            "Source URL + Platform Item ID Only": profile_row["source_url_plus_item_id_only"],
            "ACCEPT": statuses["ACCEPT"],
            "REJECT": statuses["REJECT"],
            "UNSURE": statuses["UNSURE"],
            "MOVE_TO_OTHER_CATEGORY": statuses["MOVE_TO_OTHER_CATEGORY"],
            "PENDING": statuses["PENDING"],
            "Estimated Precision Before Cleanup": manual_row["estimated_precision"],
            "Estimated Precision After Cleanup": ready_row["estimated_precision_after_cleanup"],
            "Estimated Accepted Yield": ready_row["estimated_accepted_yield"],
            "Health": ready_row["final_health"],
            "Readiness Group": ready_row["readiness_group"],
            "Main Issue": manual_row["main_issue"],
            "Recommendation": ready_row["recommendation"],
        })
    READINESS_PATH.parent.mkdir(parents=True, exist_ok=True)
    temporary = READINESS_PATH.with_suffix(READINESS_PATH.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(READINESS_PATH)


def write_source_counts(category_rows: dict[str, list[dict[str, str]]]) -> None:
    fields = ["Category", "Source Name", "Source URL", "Candidate Count", "Category Share"]
    output: list[dict[str, Any]] = []
    for category, rows in category_rows.items():
        counts = Counter((row.get("Source Name") or "", row.get("Source URL") or "") for row in rows)
        for (name, url), count in sorted(counts.items(), key=lambda item: (-item[1], item[0][1])):
            output.append({
                "Category": category,
                "Source Name": name,
                "Source URL": url,
                "Candidate Count": count,
                "Category Share": round(count / len(rows), 4) if rows else 0.0,
            })
    temporary = SOURCE_COUNTS_PATH.with_suffix(SOURCE_COUNTS_PATH.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(output)
    temporary.replace(SOURCE_COUNTS_PATH)


def main() -> None:
    parser = argparse.ArgumentParser(description="Rank review queues and build final consolidation workbooks.")
    parser.add_argument("--top-per-category", type=int, default=100)
    args = parser.parse_args()
    if not 1 <= args.top_per_category <= 100:
        raise SystemExit("--top-per-category must be between 1 and 100")
    taxonomy = read_json(PROJECT_ROOT / "config" / "taxonomy.json")
    categories = [item["code"] for item in taxonomy["subcategories"]]
    registry = source_registry()
    all_rows: dict[str, list[dict[str, str]]] = {}
    top_rows: dict[str, list[dict[str, str]]] = {}
    for category in categories:
        _, rows = enriched_rows(category, registry)
        all_rows[category] = rows
        pending = [row for row in rows if row.get("Review Status") == "PENDING"]
        top_rows[category] = balanced_pending(pending, args.top_per_category)
    write_workbook(MASTER_PATH, all_rows)
    write_workbook(TOP_PATH, top_rows, top=True)
    write_readiness(
        all_rows,
        read_json(PROFILE_PATH),
        read_json(BASELINE_PATH),
        read_json(MANUAL_PATH),
        read_json(READINESS_CONFIG_PATH),
    )
    write_source_counts(all_rows)
    print(f"MASTER\t{MASTER_PATH}")
    print(f"TOP\t{TOP_PATH}")
    print(f"READINESS\t{READINESS_PATH}")
    print(f"SOURCE_COUNTS\t{SOURCE_COUNTS_PATH}")
    print(f"CANDIDATES\t{sum(len(rows) for rows in all_rows.values())}")


if __name__ == "__main__":
    main()
